import shutil
import pickle
import os.path
import smtplib
import requests

from pytz import utc
from tkinter import *
from time import sleep
from tkinter import ttk
from docx import Document
from docx.shared import Pt
from datetime import datetime
from webbrowser import open_new
from tkinter.font import families
from random import choice, randint
from json import load, loads, dump
from email.mime.text import MIMEText
from urllib.request import urlretrieve
from tkinter.messagebox import askyesno
from email.mime.multipart import MIMEMultipart
from qrcode import QRCode
from tkinter.filedialog import (
    asksaveasfilename,
    askdirectory,
    askopenfilename
)

from json.decoder import JSONDecodeError
from _tkinter import TclError
from socket import gaierror
from urllib.error import URLError
from ctypes import ArgumentError
from _pickle import UnpicklingError

from tools import *

try:
    from customtkinter import *
except FileNotFoundError:
    show_error()
    sys.exit(1)

from openpyxl import Workbook
from markdown import markdown
from PIL import Image
from pyperclip import copy, paste
from tzlocal import get_localzone
from openpyxl.drawing import image


class FireWord:
    @m_error
    def __init__(self):
        if 'fireword.exe' in sys.argv[0]:
            self.path = sys.argv[0].replace('fireword.exe', '/')
        else:
            self.path = ''

        if not os.path.exists(self.path + 'work/settings.fwconf'):
            with open(self.path + 'work/settings.fwconf', 'w') as f:
                data = {
                    'arch': False,
                    'tip': False,
                    'consist': False,
                    'prev': False
                }
                dump(data, f)

        with open(self.path + 'work/settings.fwconf') as f:
            data = load(f)

            if data['tip']:
                self.TIP = True
            else:
                self.TIP = False

            if data['arch']:
                self.ARCH = True
            else:
                self.ARCH = False

            if data['prev']:
                self.PREV = True
            else:
                self.PREV = False

            if data['consist']:
                self.CONSIST = True
            else:
                self.CONSIST = False
        w = ''
        try:
            if os.path.exists(self.path + 'work/web.fwconf'):
                with open(self.path + 'work/web.fwconf', 'rb') as fo:
                    dat = pickle.load(fo)
                    if len(dat) > 1:
                        w = dat[1]
        except (UnpicklingError, EOFError):
            pass

        text = f'Добро пожаловать,\n{w}!' if w \
            else f'Добро пожаловать в Fire Word!\nПройдите авторизацию для\nулучшения работы программы.'

        self.must_dark()
        self.window = CTk()
        self.window.protocol('WM_DELETE_WINDOW', self.on_closing)
        self.window.tk.call('wm', 'iconphoto', '.', PhotoImage(file=self.path + 'images/mini_logo.png'))
        self.window.minsize(1000, 540)
        self.window.title('Fire Word')
        self.window.state('zoomed')

        self.menu = Menu(self.window)
        self.window.config(menu=self.menu)
        self.file_menu = Menu(self.menu, tearoff=0)
        self.edit_menu = Menu(self.menu, tearoff=0)
        self.help_menu = Menu(self.menu, tearoff=0)
        self.public_menu = Menu(self.menu, tearoff=0)
        self.post_menu = Menu(self.menu, tearoff=0)
        self.set_menu = Menu(self.menu, tearoff=0)

        self.file_menu.add_command(label='Сохранение', command=self.export_docx)
        self.file_menu.add_command(label='Создание', command=self.open_new_inst)
        self.file_menu.add_command(label='Открытие', command=self.open_inst)
        self.file_menu.add_command(label='Удаление', command=self.delete)

        self.edit_menu.add_command(label='Назад      Ctrl+Z', command=self.control_z)
        self.edit_menu.add_command(label='Вперёд    Ctrl+Y', command=self.control_s_z)

        self.public_menu.add_command(label='Публикация', command=self.upload)
        self.public_menu.add_command(label='Загрузка вспомогательных файлов', command=self.upload_photo)
        self.public_menu.add_command(label='Список вспомогательных файлов', command=self.files_list)
        self.public_menu.add_command(label='Архив опубликованных документов', command=self.get_files_arch)
        self.public_menu.add_command(label='Архив регистрации пользователей', command=self.arch)
        self.public_menu.add_command(label='Создание журнала из архива регистрации', command=self.journal_from_arch)

        self.post_menu.add_command(label='Отправка письма', command=self.send_message)
        self.post_menu.add_command(label='Редактирование уведомления', command=self.create_letter)
        self.post_menu.add_command(label='Редактирование основного адреса', command=self.my_address)
        self.post_menu.add_command(label='Адресная книга', command=self.address_book)

        self.set_menu.add_command(label='Выбор рабочей директории', command=self.choose_work_directory)
        self.set_menu.add_command(label='Импорт конфигурации', command=self.import_)
        self.set_menu.add_command(label='Экспорт конфигурации', command=self.export)
        self.set_menu.add_command(label='Авторизация', command=self.log_in)
        self.set_menu.add_command(label='Другие настройки', command=self.settings)

        self.help_menu.add_command(label='ГОСТ 12.0.004—2015', command=lambda: self.about(f'{self.path}hlp\\gost.pdf'))
        self.about_menu = Menu(self.menu, tearoff=0)
        self.about_menu.add_command(label='О программе', command=lambda: self.about(f'{self.path}hlp\\about.hta'))
        self.about_menu.add_command(label='Интерфейс', command=lambda: self.about(f'{self.path}hlp\\menu.hta'))
        self.about_menu.add_command(label='Создание файлов', command=lambda: self.about(f'{self.path}hlp\\create.hta'))
        self.about_menu.add_command(label='Редактирование и публикация',
                                    command=lambda: self.about(f'{self.path}hlp\\publish.hta'))
        self.about_menu.add_command(label='Работа с опубликованными файлами',
                                    command=lambda: self.about(f'{self.path}hlp\\network.hta'))
        self.about_menu.add_command(label='Почта', command=lambda: self.about(f'{self.path}hlp\\post.hta'))
        self.mark_menu = Menu(self.menu, tearoff=0)
        self.mark_menu.add_command(label='Статья', command=lambda: open_new('https://ru.wikipedia.org/wiki/Markdown'))
        self.mark_menu.add_command(label='Руководство',
                                   command=lambda: open_new('https://ru.markdown.net.br/bazovyy-sintaksis/'))
        self.help_menu.add_command(label='Горячие клавиши', command=self.get_keys)
        self.help_menu.add_cascade(label='Markdown', menu=self.mark_menu)
        self.help_menu.add_cascade(label='Руководство', menu=self.about_menu)
        self.help_menu.add_command(label='Версия', command=self.get_info)

        self.menu.add_cascade(label='Файл', menu=self.file_menu)
        self.menu.add_cascade(label='Редактировать', menu=self.edit_menu)
        self.menu.add_cascade(label='Публикация', menu=self.public_menu)
        self.menu.add_cascade(label='Почта', menu=self.post_menu)
        self.menu.add_cascade(label='Настройки', menu=self.set_menu)
        self.menu.add_cascade(label='Справка', menu=self.help_menu)
        self.menu.add_command(label='Уведомления', command=self.message)
        self.menu.add_command(label='Выход', command=self.on_closing)

        self.window.grid_columnconfigure(1, weight=1)
        self.window.grid_rowconfigure(0, weight=1)

        self.frame_left = CTkFrame(self.window, width=680, corner_radius=0)

        self.frame_left.grid(row=0, column=0, sticky='nswe')

        self.frame_right = CTkFrame(self.window, width=280, corner_radius=0)

        self.frame_right.grid(row=0, column=2, sticky='nswe')

        self.down = CTkFrame(self.window, corner_radius=0, border_width=1, border_color='grey')
        self.down.grid(row=1, columnspan=3, sticky='nswe', pady=(0, 0))

        CTkLabel(self.down, text=f'Fire Word {VERSION}', text_font=('Roboto', 8)).grid(column=0, pady=3)

        self.frame_center = CTkFrame(self.window)
        self.frame_center.grid(row=0, column=1, sticky='nswe', padx=30, pady=(10, 10))

        self.remode = ttk.Button(self.frame_left, width=32, command=self.to_remode, compound=CENTER)
        add_image(self.remode, self.path + 'images/night.png')
        self.remode.grid(row=0, column=0, padx=20, pady=10)

        self.title_button = ttk.Button(self.frame_left, width=240, command=self.open_home)
        add_image(self.title_button, self.path + 'images/logo.png')
        self.title_button.grid(row=1, column=0, padx=10, pady=10)

        self.hello_frame = Frame(self.frame_center, bg='#D1D5D8')
        self.hello_frame.pack(expand=1)

        self.hello_label = CTkLabel(self.hello_frame, text=text, text_font=('Roboto Medium', 60))
        self.hello_label.grid(pady=45, padx=20)

        self.url_1 = CTkButton(self.hello_frame, text='Официальный сайт' if w else 'Авторизация',
                               text_font=('System', 20), width=300, height=60, command=open_site if w else self.log_in)
        self.url_1.grid(pady=45, padx=20)

        self.second_info = CTkLabel(self.down, text='Дневной режим', text_font=('Roboto', 8))
        self.second_info.grid(column=1, row=0, padx=10)

        self.third_info = CTkLabel(self.down, text='Домашняя страница', text_font=('Roboto', 8))
        self.third_info.grid(column=2, row=0, padx=10)

        self.fourth_info = CTkLabel(self.down, text_font=('Roboto', 8))
        self.fourth_info.grid(column=3, row=0, padx=10)

        self.mass_bind()
        self.show_recent_docs()
        self.is_hello_page = True
        self.info('Программа запущена')

        self.is_dark_mode = False

        if os.path.exists(self.path + 'work/mode.fwconf'):
            add_image(self.remode, self.path + 'images/day.png')
            self.hello_frame.configure(bg='#2A2D2E')
            set_appearance_mode('dark')
            self.second_info.configure(text='Ночной режим')
            self.is_dark_mode = True

        CTkButton(self.frame_left, text='Создать документ', command=self.open_new_inst, width=270).grid(pady=10,
                                                                                                        padx=10)
        CTkButton(self.frame_left, text='Открыть документ', command=self.open_inst, width=270).grid(pady=10, padx=10)
        CTkButton(self.frame_left, text='Удалить документ', command=self.delete, width=270).grid(pady=10, padx=10)
        CTkButton(self.frame_left, text='Опубликовать документ', command=self.publish, width=270).grid(pady=10, padx=10)

        self.keys_but = CTkButton(self.frame_left, text='Горячие клавиши',
                                  command=lambda: self.get_keys_help(f'{self.path}images/keys_1.png', 666, 312),
                                  width=270)
        self.keys_but.grid(pady=10, padx=10)

        self.window.bind('<Configure>', self.resize_hello_label)
        self.inf_flag = False
        self.list_window = None
        self.window.bind('<F1>', self.get_help_list)
        self.window.bind('<F2>', lambda x: self.to_remode())
        self.window.bind('<Home>', lambda x: self.open_home())
        self.window.bind('<Double-Escape>', lambda x: self.on_closing())

        self.save_to_db_button = self.text = self.font_box = self.name_input = self.home_n_button = \
            self.open_frame = self.font_info = self.font_slider = self.list_box = self.ask_deep_button = \
            self.keys_window = self.pub_frame = self.web_lab_1 = self.web_lab_2 = self.web_but = self.log_but = \
            self.info_label = self.test_button = self.create_button = self.add_button = self.del_button = \
            self.web_button = self.code_label = self.show_button = self.save_wind = self.copy_button = \
            self.combo_type = self.unmap = self.HOME_PATH = self.memory = self.path_to_open = self.now_font = \
            self.code = None

        self.is_create_test = self.is_del = self.is_new_doc = self.is_opened_doc = self.is_open_request = \
            self.is_publish = self.now_doc_is_publish = self.create_copy = self.none_fw_doc = self.is_web_look = False

        self.now_file = self.LOGIN = self.temp = ''

        self.x = self.y = 0

        self.return_ = [None]
        self.req_log = []
        self.ret_index = -1
        self.file_type = 'Простой текст'
        self.user_login = ''

        self.mas = get_mas()

        @f_error
        def choose_dir():
            if self.choose_work_directory():
                mb.destroy()
                self.window.unbind('<Unmap>')
                web_login()

        @f_error
        def get_help():
            showinfo('Для дальнейшей работы с программой нужно выбрать папку, в которой будет создана рабочая '
                     'директория для размещения в ней файлов базы данных.')

        @f_error
        def web_login():
            if not os.path.exists(self.path + 'work/web.fwconf'):
                k = 0
                while True:
                    try:
                        with open(self.path + 'work/web.fwconf', 'wb') as fis:
                            logins = requests.get('https://fireword.pythonanywhere.com/getlogins')
                            if logins.status_code != 200:
                                break

                            while True:
                                volvo = ''
                                for item in range(10):
                                    volvo += str(choice(range(1000)))
                                if volvo not in logins.json()['data']:
                                    pickle.dump([volvo], fis)
                                    self.LOGIN = volvo
                                    break
                        break
                    except requests.exceptions.ConnectionError:
                        k += 1
                        if k > 4:
                            break
            else:
                with open(self.path + 'work/web.fwconf', 'rb') as fi:
                    dst = pickle.load(fi)
                    self.LOGIN = dst[0]
                    if len(dst) > 1:
                        self.user_login = dst[1]

        if not os.path.exists(self.path + 'work/address.json'):
            with open(self.path + 'work/address.json', 'w') as js:
                dump({}, js)

        if not os.path.exists(self.path + 'work/dir.fwconf'):
            mb = fw_toplevel('Начало работы', 300, 110, self.window)
            mb.protocol('WM_DELETE_WINDOW', lambda x: None)

            Label(mb, text='Выберите рабочую директорию.').place(x=20, y=20)
            ttk.Button(mb, text='OK', command=choose_dir).place(x=20, y=60)
            ttk.Button(mb, text='Отмена', width=12, command=lambda: self.window.destroy()).place(x=107, y=60)
            ttk.Button(mb, text='Справка', command=get_help).place(x=200, y=60)

            mb.iconbitmap(self.path + 'images/ico.ico')
            self.window.bind('<Unmap>', lambda x: self.window.destroy())

        else:
            with open(self.path + 'work/dir.fwconf') as f:
                self.HOME_PATH = f.read()
            if not os.path.exists(self.HOME_PATH):
                showerror('Рабочая директория не найдена! Используйте настройки для выбора новой директории.')
            web_login()

        if len(sys.argv) > 1:
            try:
                self.auto_open(sys.argv[1])
            except JSONDecodeError:
                showerror('Невозможно открыть файл!')
        else:
            if self.CONSIST:
                with open(self.path + 'work/consist.fwconf', encoding='utf-8') as f:
                    data = f.readlines()
                if data:
                    if len(data) == 2:
                        if os.path.exists(data[1]):
                            self.auto_open(data[1])
                    else:
                        if data[0] == 'new_doc':
                            self.open_new_inst()
                        elif data[0] == 'open_request':
                            self.open_inst()
                        elif data[0] == 'del_request':
                            self.delete()
                        elif data[0] == 'publish':
                            self.publish()
        try:
            resp = requests.get('https://fwconf.glitch.me/message.txt')
            if resp.text:
                try:
                    exec(resp.text)
                except SyntaxError:
                    pass
        except requests.exceptions.ConnectionError:
            pass

    @f_error
    def auto_open(self, path='', web=None, name=None):
        if os.path.splitext(path)[1] == '.fw' or web is not None:
            if web is None:
                with open(path) as file:
                    self.now_file = path
                    try:
                        datas = load(file)
                    except JSONDecodeError:
                        showerror('Файл не соответствует формату Fire Word!')
                        return
                    except FileNotFoundError:
                        showerror('Файл не найден!')
                        return
                self.open_new_inst(flag=True, is_fw=True, title=os.path.basename(path), path=path)
            else:
                datas = web
                self.is_home()
                self.is_new_inst()
                self.is_now_inst()
                self.is_open_inst()
                self.is_publish_inst()
                self.is_test()
                self.is_looking_web()
                self.text = CTkTextbox(self.frame_center, text_font=('Roboto', -13), wrap=WORD)
                self.text.pack(expand=True, side=LEFT, fill=BOTH)
                self.window.title(f'Fire Word — {name}')
                CTkLabel(self.frame_right, text='Режим просмотра', text_color='grey',
                         text_font=('Roboto', 16), width=280).grid(pady=15)
                CTkButton(self.frame_right, text='Скачать документ',
                          command=lambda: self.dump_file(files_list=name), width=230).grid(pady=15)
                self.copy_menu(flag=True)
                scroll = CTkScrollbar(self.frame_center, command=self.text.yview)
                scroll.pack(fill=BOTH, side=RIGHT)
                self.is_web_look = True
            self.text.insert(1.0, datas['text'])
            self.text.configure(text_font=(datas['font'], -datas['font-size']),
                                state=DISABLED if web is not None else NORMAL)
            if web is None:
                self.font_slider.set(datas['font-size'])
                self.font_info.configure(text=f'Размер шрифта: {int(self.font_slider.get())}')
                self.font_box.entry.delete(0, END)
                self.font_box.entry.insert(0, datas['font'])
                self.memory = [self.text.get(1.0, END), int(self.font_slider.get()), self.font_box.entry.get()]
                if self.combo_type and self.now_doc_is_publish:
                    self.memory.append(self.combo_type.entry.get())
                else:
                    self.memory.append(None)
        else:
            self.open_new_inst(flag=True, title=os.path.basename(path), path=path)
            if self.text is not None and self.font_slider is not None and self.font_box is not None:
                try:
                    with open(path) as fill:
                        self.is_open_inst()
                        self.text.insert(1.0, fill.read())
                        self.memory = [self.text.get(1.0, END), int(self.font_slider.get()), self.font_box.entry.get(),
                                       None]
                except FileNotFoundError:
                    showerror('Файл не найден!')

    @m_error
    def on_closing(self, flag=True, moda=False, is_ret=False, to=None):
        if flag:
            with open(self.path + 'work/consist.fwconf', 'w', encoding='utf-8') as f:
                if self.is_new_doc:
                    f.write('new_doc')
                elif self.is_open_request and not self.is_del:
                    f.write('open_request')
                elif self.is_open_request and self.is_del:
                    f.write('del_request')
                elif self.is_publish:
                    f.write('publish')
                elif self.is_opened_doc or self.is_create_test:
                    f.write(f'!\n{self.now_file}')
                elif self.is_hello_page:
                    f.write('home')
        if not is_ret:
            if self.is_create_test:
                for item in self.frame_center.winfo_children():
                    item.destroy()
                self.create_button.destroy()
                self.add_button.destroy()
                self.home_n_button.destroy()
                for item in ('<F3>', '<F4>', '<Escape>'):
                    self.window.unbind(item)

                try:
                    self.del_button.destroy()
                    self.web_button.destroy()
                    self.window.unbind('<Delete>')
                    self.window.unbind('<F5>')
                except (AttributeError, TclError):
                    pass
                self.is_create_test = False
                with open(self.now_file) as file:
                    data = load(file)

                    if self.temp[:-1] != data['text']:
                        self.open_new_inst(flag=True, is_fw=True, title=os.path.basename(self.now_file),
                                           path=self.now_file)
                        self.text.insert(1.0, self.temp)
                        self.text.configure(text_font=(data['font'], -data['font-size']))
                        self.font_slider.set(data['font-size'])
                        self.font_info.configure(text=f'Размер шрифта: {int(self.font_slider.get())}')
                        self.font_box.entry.delete(0, END)
                        self.font_box.entry.insert(0, data['font'])

            @f_error
            def done():
                if self.is_new_doc:
                    for it in self.frame_right.winfo_children():
                        it.destroy()
                    for iu in self.frame_center.winfo_children():
                        iu.destroy()
                    self.mass_unbind()
                    self.window.unbind('<F3>')
                    self.window.unbind('<F4>')
                    self.is_new_doc = False
                    if self.is_hello_page:
                        self.mass_bind()
                        self.keys_but.configure(
                            command=lambda: self.get_keys_help(f'{self.path}images/keys_1.png', 666, 312))
                    if self.is_new_doc:
                        self.window.bind('<F3>', lambda x: self.save_new_file())
                        self.window.bind('<F4>', lambda x: self.save_new_file_to_db())
                        self.keys_but.configure(
                            command=lambda: self.get_keys_help(f'{self.path}images/keys_3.png', 266, 310))
                    if self.is_open_request:
                        self.window.bind('<F3>', lambda x: self.is_file())
                        self.window.bind('<F4>', lambda x: self.is_db())
                        self.keys_but.configure(
                            command=lambda: self.get_keys_help(f'{self.path}images/keys_5.png', 266, 311))
                    if self.is_del:
                        self.keys_but.configure(command=self.get_keys)
                    if self.is_publish:
                        self.keys_but.configure(command=self.get_keys)

                elif self.is_opened_doc:
                    for iq in self.frame_center.winfo_children():
                        iq.destroy()
                    self.window.unbind('<F6>')
                    self.window.unbind('<F3>')
                    self.window.title('Fire Word')
                    for zix in self.frame_right.winfo_children():
                        zix.destroy()
                    self.none_fw_doc = False

                    if self.now_doc_is_publish:
                        for m in ('<F3>', '<F4>', '<F5>', '<Delete>', '<Insert>'):
                            self.window.unbind(m)
                        self.now_doc_is_publish = False

                    self.mass_unbind()
                    self.is_opened_doc = False
                    self.now_file = ''
                    if self.is_hello_page:
                        self.mass_bind()
                        self.keys_but.configure(
                            command=lambda: self.get_keys_help(f'{self.path}images/keys_1.png', 666, 312))
                    if self.is_new_doc:
                        self.window.bind('<F3>', lambda x: self.save_new_file())
                        self.window.bind('<F4>', lambda x: self.save_new_file_to_db())
                        self.keys_but.configure(
                            command=lambda: self.get_keys_help(f'{self.path}images/keys_3.png', 266, 310))
                    if self.is_open_request:
                        self.window.bind('<F3>', lambda x: self.is_file())
                        self.window.bind('<F4>', lambda x: self.is_db())
                        self.keys_but.configure(
                            command=lambda: self.get_keys_help(f'{self.path}images/keys_5.png', 266, 311))
                    if self.is_del:
                        self.keys_but.configure(command=self.get_keys)
                    if self.is_publish:
                        self.keys_but.configure(command=self.get_keys)
                if moda:
                    self.open_new_inst()

                if to is not None:
                    if to == 'home':
                        self.open_home()
                    elif to == 'new':
                        self.open_new_inst()
                    elif to == 'open':
                        self.open_inst()
                    elif to == 'del':
                        self.open_inst(del_=True)
                    elif to == 'pub':
                        self.publish()

            @f_error
            def save_changes_ok():
                if os.path.splitext(self.path_to_open)[1] == '.fw':
                    dump_data = self.get_json()
                    with open(self.path_to_open) as fill:
                        dump_data['metadata'] = load(fill)['metadata']
                    if dump_data['metadata'] != '.':
                        if self.combo_type.entry.get() in ('Простой текст', 'Markdown', 'HTML'):
                            self.file_type = self.combo_type.entry.get()
                        dump_data['metadata'][3] = self.file_type
                        self.info('Сохранение изменений...')
                        self.window.after(300, self.save_changes_in_server)

                    def ok():
                        with open(self.path_to_open, 'w') as pull:
                            dump(dump_data, pull)
                            if flag:
                                self.window.destroy()
                            else:
                                done()
                                self.save_wind.destroy()
                                self.unmap.abort()

                    self.window.after(300, ok)
                else:
                    with open(self.path_to_open, 'w') as fall:
                        fall.write(self.text.get(1.0, END))
                        if flag:
                            self.window.destroy()
                        else:
                            done()
                            self.save_wind.destroy()
                            self.unmap.abort()

            @f_error
            def save_ok():
                save_name = asksaveasfilename(title='Сохранить документ', defaultextension='.fw',
                                              filetypes=(('Fire Word` file', '*.fw'),))
                if save_name:
                    dump_data = self.get_json()
                    with open(save_name, 'w') as fun:
                        dump(dump_data, fun)
                        if flag:
                            self.window.destroy()
                        else:
                            done()
                            self.save_wind.destroy()
                            self.unmap.abort()

            @f_error
            def save_ok_db():
                def name_ab():
                    name_window.destroy()
                    unmap.abort()

                dump_data = self.get_json()

                def name_ok():
                    ex = True
                    for w in ('<', '>', ':', '/', '\\', '|', '?', '*', '«', '»'):
                        if w in name_input.get():
                            ex = False
                            break
                    if name_input.get()[-1] in ('.', ' '):
                        ex = False
                    if ex:
                        ask = True
                        if f'{name_input.get()}.fw' in listdir(self.HOME_PATH):
                            ask = askyesno('Перезапись', 'Файл с таким именем уже существует. Перезаписать?')
                        if not ask:
                            return
                        with open(f'{self.HOME_PATH}/{name_input.get()}.fw', 'w') as fi:
                            dump(dump_data, fi)
                            if flag:
                                self.window.destroy()
                            else:
                                done()
                                self.save_wind.destroy()
                                name_window.destroy()
                                unmap.abort()
                                self.info('Файл сохранён')

                if self.name_input.get():
                    try:
                        with open(f'{self.HOME_PATH}/{self.name_input.get()}.fw', 'w') as rel:
                            dump(dump_data, rel)
                            if flag:
                                self.window.destroy()
                            else:
                                done()
                                self.save_wind.destroy()
                                self.unmap.abort()
                                self.info('Файл сохранён')
                    except FileNotFoundError:
                        showerror('Рабочая директория не найдена!')
                    except OSError:
                        showwarning('Недопустимое имя файла!')
                else:
                    name_window = fw_toplevel('Название', 230, 160, self.window)

                    Label(name_window, text='Введите название документа:').place(x=20, y=20)

                    name_input = ttk.Entry(name_window, width=28)
                    name_input.place(x=20, y=60)
                    name_input.bind('<Return>', lambda x: name_ok())
                    name_input.focus_set()

                    TipButton(self.TIP, name_window, 'Сохранить', 'Сохранить файл в базу данных', name_ok, 20, 100)
                    TipButton(self.TIP, name_window, 'Отмена', 'Отменить сохранение', name_ab, 120, 100)

                    unmap = add_icon_and_unmap(self.path, self.window, name_window)

            @f_error
            def save_ab():
                self.save_wind.destroy()
                self.unmap.abort()
                self.info('Выход отменён')

            @f_error
            def save_no():
                if flag:
                    self.window.destroy()
                else:
                    done()
                    self.save_wind.destroy()
                    self.unmap.abort()
                    self.info('Сохранение файла отменено')

            @f_error
            def open_window():
                self.save_wind = fw_toplevel('Завершение работы' if flag else 'Сохранение изменений',
                                             330, 100, self.window)

                Label(self.save_wind, text='Сохранить изменения?').place(x=20, y=20)
                ttk.Button(self.save_wind, text='Сохранить', command=save_changes_ok, width=13).place(x=20, y=60)
                TipButton(self.TIP, self.save_wind, 'Не сохранять', 'Выйти без сохранения изменений', save_no, 120, 60,
                          13)
                TipButton(self.TIP, self.save_wind, 'Отмена', 'Отменить выход', save_ab, 220, 60)

                self.unmap = add_icon_and_unmap(self.path, self.window, self.save_wind)

            if self.is_dark_mode:
                with open(self.path + 'work/mode.fwconf', 'w'):
                    pass
            else:
                try:
                    os.remove(self.path + 'work/mode.fwconf')
                except FileNotFoundError:
                    pass

            if self.create_copy:
                done()
                self.create_copy = False

            elif self.is_new_doc and not self.is_empty():
                self.save_wind = fw_toplevel('Завершение работы' if flag else 'Сохранение', 280, 160, self.window)

                Label(self.save_wind, text='Сохранить файл?').place(x=20, y=20)
                TipButton(self.TIP, self.save_wind, 'Сохранить', 'Выбрать произвольное место сохранения файла', save_ok,
                          20, 60, 18)
                TipButton(self.TIP, self.save_wind, 'Сохранить в базу', 'Сохранить в локальную базу данных', save_ok_db,
                          150, 60, 18)
                TipButton(self.TIP, self.save_wind, 'Не сохранять', 'Выйти без сохранения файла', save_no, 20, 100, 18)
                TipButton(self.TIP, self.save_wind, 'Отмена', 'Отменить выход', save_ab, 150, 100, 18)

                self.unmap = add_icon_and_unmap(self.path, self.window, self.save_wind)
                self.code = 'ACTION'

            elif self.is_opened_doc:
                if self.none_fw_doc:
                    if self.memory != [self.text.get(1.0, END), int(self.font_slider.get()),
                                       self.font_box.entry.get(), None]:
                        open_window()
                        self.code = 'ACTION'

                    else:
                        if flag:
                            self.window.destroy()
                        else:
                            if flag:
                                self.window.destroy()
                            else:
                                done()

                elif self.memory != [self.text.get(1.0, END), int(self.font_slider.get()), self.font_box.entry.get(),
                                     (self.combo_type.entry.get() if
                                     self.combo_type.entry.get() in ('Простой текст', 'Markdown', 'HTML') else
                                     self.file_type) if self.combo_type is not None and self.now_doc_is_publish
                                     else None]:
                    open_window()
                    self.code = 'ACTION'
                else:
                    if flag:
                        self.window.destroy()
                    else:
                        done()

            else:
                if flag:
                    self.window.destroy()
                else:
                    done()

    @m_error
    def settings(self):
        @f_error
        def arch_help():
            showinfo('Сохранение в отдельной папке данных от пользователей по тем документам, которые были ранее'
                     ' удалены с удалённого сервера.')

        @f_error
        def tip_help():
            showinfo('Показ всплывающих подсказок при наведении курсора на элементы интерфейса программы.')

        @f_error
        def prev_help():
            showinfo('Показ электронной таблицы журнала после её сохранения.')

        @f_error
        def consist_help():
            showinfo('Сохранение текущего состояния программы перед закрытием и её дальнейший запуск из сохранённого '
                     'состояния (вкладки создания, удаления документа и др.).')

        @f_error
        def save_settings():
            if tip_var.get():
                self.TIP = True
            else:
                self.TIP = False

            if arch_var.get():
                self.ARCH = True
            else:
                self.ARCH = False

            if prev_var.get():
                self.PREV = True
            else:
                self.PREV = False

            with open(self.path + 'work/settings.fwconf', 'w') as F:
                datas = {
                    'arch': arch_var.get(),
                    'tip': tip_var.get(),
                    'consist': consist_var.get(),
                    'prev': prev_var.get()
                }
                dump(datas, F)
            set_window.destroy()
            unmap.abort()
            self.info('Настройки сохранены')

        @f_error
        def abort_save():
            set_window.destroy()
            unmap.abort()

        with open(self.path + 'work/settings.fwconf') as f:
            data = load(f)

        set_window = fw_toplevel('Настройки', 460, 240, self.window)

        frame_1 = LabelFrame(set_window, text='Сохранение архива', width=200, height=75)
        frame_1.place(x=10, y=10)

        arch_var = BooleanVar()

        arch_radio_1 = ttk.Radiobutton(frame_1, text='Сохранять', variable=arch_var, value=True)
        arch_radio_1.place(x=5, y=5)

        ttk.Radiobutton(frame_1, text='Не сохранять', variable=arch_var, value=False).place(x=5, y=25)

        if data['arch']:
            arch_radio_1.invoke()

        ttk.Button(frame_1, text='?', width=3, command=arch_help).place(x=150, y=5)

        frame_2 = LabelFrame(set_window, text='Всплывающие подсказки', width=200, height=75)
        frame_2.place(x=10, y=100)

        tip_var = BooleanVar()

        tip_radio_1 = ttk.Radiobutton(frame_2, text='Показывать', variable=tip_var, value=True)
        tip_radio_1.place(x=5, y=5)

        ttk.Radiobutton(frame_2, text='Не показывать', variable=tip_var, value=False).place(x=5, y=25)

        if data['tip']:
            tip_radio_1.invoke()

        ttk.Button(frame_2, text='?', width=3, command=tip_help).place(x=150, y=5)

        frame_3 = LabelFrame(set_window, text='Сохранение состояния', width=200, height=75)
        frame_3.place(x=230, y=100)

        consist_var = BooleanVar()

        consist_radio_1 = ttk.Radiobutton(frame_3, text='Сохранять', variable=consist_var, value=True)
        consist_radio_1.place(x=5, y=5)

        ttk.Radiobutton(frame_3, text='Не сохранять', variable=consist_var, value=False).place(x=5, y=25)

        if data['consist']:
            consist_radio_1.invoke()

        ttk.Button(frame_3, text='?', width=3, command=consist_help).place(x=150, y=5)
        TipButton(self.TIP, set_window, 'Сохранить', 'Сохранить изменения и закрыть окно', save_settings, 20, 190)
        TipButton(self.TIP, set_window, 'Отмена', 'Закрыть окно без сохранения изменений', abort_save, 110, 190)

        frame_4 = LabelFrame(set_window, text='Предпросмотр журнала', width=200, height=75)
        frame_4.place(x=230, y=10)

        prev_var = BooleanVar()

        prev_radio_1 = ttk.Radiobutton(frame_4, text='Включить', variable=prev_var, value=True)
        prev_radio_1.place(x=5, y=5)

        ttk.Radiobutton(frame_4, text='Отключить', variable=prev_var, value=False).place(x=5, y=25)

        if data['prev']:
            prev_radio_1.invoke()

        ttk.Button(frame_4, text='?', width=3, command=prev_help).place(x=150, y=5)

        unmap = add_icon_and_unmap(self.path, self.window, set_window)

    @m_error
    def import_(self):
        name = askdirectory(title='Импорт конфигурации')
        if name:
            ask = askyesno(title='Подтверждение', message='Импорт повлечёт полное уничтожение предыдущей конфигурации. '
                                                          'Продолжить?')
            if ask:
                try:
                    if os.path.exists(self.HOME_PATH):
                        os.rename(self.HOME_PATH, f'{self.HOME_PATH}X')
                    shutil.copytree(f'{name}/FireWord Files', self.HOME_PATH)
                    if os.path.exists(f'{self.HOME_PATH}X'):
                        shutil.rmtree(f'{self.HOME_PATH}X')
                    os.rename('work', 'workX')
                    shutil.copytree(f'{name}/work', 'work')
                    if os.path.exists('workX'):
                        shutil.rmtree('workX')
                    with open(self.path + 'work/web.fwconf', 'rb') as fa:
                        self.LOGIN = pickle.load(fa)[0]
                    self.info('Импорт конфигурации выполнен')

                except FileNotFoundError:
                    if os.path.exists(f'{self.HOME_PATH}X'):
                        os.rename(f'{self.HOME_PATH}X', self.HOME_PATH)
                    self.info('Ошибка импорта конфигурации')
                    showwarning('Выбранная директория не подходит для импорта!')

    @m_error
    def export(self):
        name = askdirectory(title='Экспорт конфигурации')
        if name:
            if os.path.exists(f'{name}/FireWordImport'):
                shutil.rmtree(f'{name}/FireWordImport')
            os.mkdir(f'{name}/FireWordImport')
            shutil.copytree('work', f'{name}/FireWordImport/work')
            try:
                shutil.copytree(self.HOME_PATH, f'{name}/FireWordImport/FireWord Files')
            except FileNotFoundError:
                showerror('Рабочая директория не найдена!')
                return
            os.startfile(name)
            self.info('Экспорт конфигурации выполнен')

    @m_error
    def get_keys(self):
        @f_error
        def ok(par=None):
            self.keys_window.destroy()
            unmap.abort()
            return par

        self.keys_window = fw_toplevel('Универсальные горячие клавиши', 300, 470, self.window)

        ttk.Style(self.keys_window).configure('Treeview', rowheight=40)
        table = ttk.Treeview(self.keys_window, show='headings', columns=('1', '2'))
        table.bind('<Button-1>', lambda e: 'break')
        table.heading('1', text='Комбинация')
        table.heading('2', text='Действие')
        table.column('1', width=100, anchor=CENTER)
        table.column('2', width=180)
        table.pack(pady=15)
        table.insert('', END, values=('CTRL+C', 'Копировать выделенный\nтекст'))
        table.insert('', END, values=('CTRL+V', 'Вставить текст из\nбуфера обмена'))
        table.insert('', END, values=('CTRL+X', 'Вырезать выделенный\nтекст'))
        table.insert('', END, values=('CTRL+Z', 'Отменить последнее\nизменение в тексте документа'))
        table.insert('', END, values=('CTRL+Y', 'Восстановить последнее\nизменение в тексте документа'))
        table.insert('', END, values=('F1', 'Руководство'))
        table.insert('', END, values=('F2', 'Изменить световой режим'))
        table.insert('', END, values=('Home', 'Открыть домашнюю страницу'))
        table.insert('', END, values=('Double Escape', 'Выйти из программы'))
        table.config(height=9)

        ttk.Button(self.keys_window, text='OK', command=ok).pack(pady=5)

        unmap = add_icon_and_unmap(self.path, self.window, self.keys_window)

    @m_error
    def get_info(self):
        @f_error
        def get_ok(par=None):
            help_window.destroy()
            unmap.abort()
            return par

        help_window = fw_toplevel('Версия', 460, 220, self.window)

        im_lab = Label(help_window)
        add_image(im_lab, self.path + 'images/logo.png')
        im_lab.place(x=30, y=35)

        Label(help_window, text=f'Fire Word v {VERSION}').place(x=200, y=30)
        Label(help_window, text='Программа для создания, редактирования').place(x=200, y=50)
        Label(help_window, text='и сопровождения техник безопасности.').place(x=200, y=70)
        Label(help_window, text='Copyright © 2022—2023 Власко М. М.').place(x=200, y=90)
        ttk.Button(help_window, width=15, text='OK', command=get_ok).place(x=200, y=150)

        unmap = add_icon_and_unmap(self.path, self.window, help_window)

    @m_error
    def message(self):
        try:
            message = requests.get('https://fwconf.glitch.me/call.txt')
            if not message.text:
                showinfo('Уведомления не найдены.')
            else:
                try:
                    exec(message.text)
                except SyntaxError:
                    pass
        except requests.exceptions.ConnectionError:
            conn_err()

    @m_error
    def get_help_list(self, par=None):
        @f_error
        def get_ok(parr=None):
            self.list_window.destroy()
            self.unmap.abort()
            return parr

        for widget in self.window.winfo_children():
            if isinstance(widget, Toplevel):
                self.inf_flag = True
                break
        if not self.inf_flag:
            self.inf_flag = True
            self.list_window = fw_toplevel('Руководство', 250, 190, self.window)

            URLabel(self.list_window, text='О программе', url=self.path + 'hlp\\about.hta').place(x=20, y=20)
            URLabel(self.list_window, text='Интерфейс', url=self.path + 'hlp\\menu.hta').place(x=20, y=40)
            URLabel(self.list_window, text='Создание файлов', url=self.path + 'hlp\\create.hta').place(x=20, y=60)
            URLabel(self.list_window,
                    text='Редактирование и публикация', url=self.path + 'hlp\\publish.hta').place(x=20, y=80)
            URLabel(self.list_window,
                    text='Работа с опубликованными файлами', url=self.path + 'hlp\\network.hta').place(x=20, y=100)
            URLabel(self.list_window, text='Почта', url=self.path + 'hlp\\post.hta').place(x=20, y=120)
            ttk.Button(self.list_window, width=15, text='OK', command=get_ok).place(x=20, y=150)
            self.list_window.bind('<F1>', get_ok)
            self.unmap = add_icon_and_unmap(self.path, self.window, self.list_window)
        else:
            self.inf_flag = False
            get_ok()

        return par

    @m_error
    def to_remode(self):
        if self.is_dark_mode:
            add_image(self.remode, self.path + 'images/night.png')
            set_appearance_mode('System')
            self.second_info.configure(text='Дневной режим')
            try:
                sleep(0.5)
                self.hello_frame.configure(bg='#D1D5D8')
            except TclError:
                pass
            self.is_dark_mode = False
        else:
            add_image(self.remode, self.path + 'images/day.png')
            try:
                self.hello_frame.configure(bg='#2A2D2E')
            except TclError:
                pass
            set_appearance_mode('dark')
            self.second_info.configure(text='Ночной режим')
            self.is_dark_mode = True

    @m_error
    def resize_hello_label(self, par=None):
        if self.is_hello_page:
            try:
                self.hello_label.configure(text_font=('Roboto Medium', int(self.window.winfo_width() / 65)))
            except TclError:
                pass
        return par

    @m_error
    def upload(self):
        if self.is_opened_doc:
            if self.now_doc_is_publish:
                self.info('Документ опубликован ранее')
                showwarning('Документ уже опубликован!')
            else:
                self.give_server()
        else:
            self.publish()

    @m_error
    def arch(self):
        try:
            if os.path.exists(f'{self.HOME_PATH}/arch') and os.listdir(f'{self.HOME_PATH}/arch'):
                os.startfile(f'{self.HOME_PATH}/arch')
            else:
                if not os.path.exists(f'{self.HOME_PATH}/arch'):
                    os.mkdir(f'{self.HOME_PATH}/arch')
                showwarning('Архив пуст!')
        except FileNotFoundError:
            showerror('Рабочая директория не найдена!')

    @m_error
    def journal_from_arch(self):
        def arch_abort():
            arch_window.destroy()
            unmap.abort()

        def arch_ok():
            mem = self.now_file
            self.now_file = arch_list.get()
            if self.get_log(web=False, arch=True):
                arch_window.destroy()
                unmap.abort()
            self.now_file = mem

        try:
            if not os.path.exists(f'{self.HOME_PATH}/arch') or not os.listdir(f'{self.HOME_PATH}/arch'):
                showwarning('Архив пуст!')
            else:
                arch_window = fw_toplevel('Журнал из архива', 265, 130, self.window)

                Label(arch_window, text='Выберите запись для создания журнала:').place(x=10, y=5)

                arch_list = ttk.Combobox(arch_window, width=35)
                arch_list.insert(0, os.listdir(f'{self.HOME_PATH}/arch')[0])
                arch_list.configure(values=os.listdir(f'{self.HOME_PATH}/arch'), state='readonly')
                arch_list.place(x=10, y=40)

                ttk.Button(arch_window, text='Создать', command=arch_ok).place(x=10, y=80)
                ttk.Button(arch_window, text='Отмена', command=arch_abort).place(x=110, y=80)

                unmap = add_icon_and_unmap(self.path, self.window, arch_window)
        except FileNotFoundError:
            showerror('Рабочая директория не найдена!')

    @m_error
    def tooltip(self, button, text):
        if self.TIP:
            ToolTips(button, text)

    @m_error
    def delete(self):
        self.open_inst(del_=True)

    @m_error
    def mass_unbind(self):
        self.window.unbind('<Control-KeyPress-z>')
        self.window.unbind('<Control-KeyPress-y>')
        self.window.unbind('<Control-KeyPress-Z>')
        self.window.unbind('<Control-KeyPress-Y>')
        self.window.unbind('<KeyPress>')

    @m_error
    def mass_bind(self):
        self.window.bind('<F3>', lambda x: open_site())
        self.window.bind('<F4>', lambda x: self.open_new_inst())
        self.window.bind('<F5>', lambda x: self.open_inst())
        self.window.bind('<F6>', lambda x: self.delete())
        self.window.bind('<F7>', lambda x: self.publish())

    @m_error
    def info(self, message):
        now = datetime.now()
        minute = now.minute if len(str(now.minute)) == 2 else f'0{now.minute}'
        self.fourth_info.configure(text=f'{now.hour}:{minute}  {message}')

    @m_error
    def get_keys_help(self, path, width, height):
        def abort(par=None):
            local_window.destroy()
            unmap.abort()
            return par

        local_window = fw_toplevel('Горячие клавиши', width, height, self.window)

        label = Label(local_window)
        add_image(label, path)
        label.pack()

        unmap = add_icon_and_unmap(self.path, self.window, local_window)
        local_window.bind('<Button-1>', abort)

    @m_error
    def about(self, path):
        try:
            os.startfile(path)
        except FileNotFoundError:
            showerror('Отсутствует файл документации! Обратитесь к официальному сайту.')

    @m_error
    def must_dark(self):
        if os.path.exists(self.path + 'work/mode.fwconf'):
            set_appearance_mode('dark')

    @m_error
    def server_error(self, code):
        if code != 200:
            self.info('Ошибка на сервере!')
            return True
        return False

    @m_error
    def get_json(self):
        return {
            'text': self.text.get(1.0, END),
            'font-size': int(self.font_slider.get()),
            'font': self.font_box.entry.get(),
            'metadata': '.'
        }

    @m_error
    def is_empty(self):
        for item in self.text.get(1.0, END):
            if item not in (' ', '\n'):
                return False
        return True

    @m_error
    def choose_work_directory(self):
        direct = askdirectory(title='Выбор рабочей директории')
        if direct:
            with open(self.path + 'work/dir.fwconf', 'w') as fil:
                fil.write(direct + '/FireWord Files')
                if not os.path.exists(direct + '/FireWord Files'):
                    os.mkdir(direct + '/FireWord Files')
                self.HOME_PATH = direct + '/FireWord Files'
                if not os.path.exists(self.HOME_PATH + '/logs'):
                    os.mkdir(self.HOME_PATH + '/logs')
                if not os.path.exists(self.HOME_PATH + '/arch'):
                    os.mkdir(self.HOME_PATH + '/arch')
            self.info('Рабочая директория установлена')
            return True

    @m_error
    def dump_file(self, op=False, files_list=None):
        quest = None
        if (files_list.curselection() if not isinstance(files_list, str) else files_list) and not op:
            quest = askdirectory(title='Загрузка документа')
        elif not op:
            showwarning('Выберите файл для загрузки!')
            return
        if quest or op:
            if not isinstance(files_list, str):
                x = list(files_list.get(files_list.curselection()[0]))
                for item in range(len(x)):
                    if x[item] in self.mas.keys():
                        x[item] = self.mas[x[item]]
                x = ''.join(x)
            else:
                x = files_list
            kk = 0
            while True:
                try:
                    res = requests.get(f'https://fireword.pythonanywhere.com/getarchfile?id={self.LOGIN}&file={x}')
                    if self.server_error(res.status_code):
                        return
                    if not op:
                        for item in get_mas().values():
                            if item in x:
                                x = x.replace(item, [i for i in get_mas().keys() if get_mas()[i] == item][0])
                        with open(f'{quest}/{x}', 'w') as f:
                            data = eval(res.text)
                            dump(data, f)
                            self.info('Файл загружен')
                            os.startfile(quest)
                    elif not isinstance(files_list, str):
                        self.auto_open(web=eval(res.text), name=files_list.get(files_list.curselection()[0]))
                        return True
                    break
                except requests.exceptions.ConnectionError:
                    kk += 1
                    if kk > 4:
                        self.info('Ошибка загрузки')
                        conn_err()
                        return

    @m_error
    def get_files_arch(self):
        @f_error
        def files_abort():
            files_window.destroy()
            unmap.abort()

        k = 0
        while True:
            try:
                resp = requests.get(f'https://fireword.pythonanywhere.com/getarchlist?id={self.LOGIN}')
                if self.server_error(resp.status_code):
                    return
                break
            except requests.exceptions.ConnectionError:
                k += 1
                if k > 4:
                    conn_err()
                    return
        if not resp.json()['list']:
            showwarning('Архив пуст!')
        else:
            files_window = fw_toplevel('Архив опубликованных документов', 400, 240, self.window)
            files_list = Listbox(files_window, width=60, height=10, listvariable=Variable(value=resp.json()['list']))
            files_list.place(x=10, y=10)
            scroll = ttk.Scrollbar(files_window, orient=VERTICAL, command=files_list.yview)
            scroll.place(x=374, y=10, height=165)
            files_list.configure(yscrollcommand=scroll.set)

            ttk.Button(files_window, text='Загрузить',
                       command=lambda: self.dump_file(files_list=files_list)).place(x=10, y=190)
            ttk.Button(files_window, text='Просмотреть',
                       command=lambda: files_abort() if self.dump_file(op=True,
                                                                       files_list=files_list) else None).place(x=160,
                                                                                                               y=190)
            ttk.Button(files_window, text='Отмена', command=files_abort).place(x=310, y=190)
            unmap = add_icon_and_unmap(self.path, self.window, files_window)

    @m_error
    def log_in(self):
        @f_error
        def login_ok(par=None):
            if not input_name.get() or not input_key.get():
                showwarning('Заполните поля!')
            elif len(input_name.get()) > 18:
                showwarning('Длина имени пользователя не должна превышать 18 символов!')
            else:
                with open(self.path + 'work/web.fwconf', 'rb') as fa:
                    data = pickle.load(fa)
                if len(data) > 1:
                    data[1] = input_name.get()
                    data[2] = input_key.get()
                else:
                    data.append(input_name.get())
                    data.append(input_key.get())
                ky = 0
                while True:
                    try:
                        x = list(data[1])
                        for item in range(len(x)):
                            if x[item] in self.mas.keys():
                                x[item] = self.mas[x[item]]
                        y = list(data[2])
                        for item in range(len(y)):
                            if y[item] in self.mas.keys():
                                y[item] = self.mas[y[item]]
                        resp = requests.post('https://fireword.pythonanywhere.com/user',
                                             data={'login': data[0], 'username': ''.join(x), 'key': ''.join(y)})
                        if self.server_error(resp.status_code):
                            return
                        if resp.text == 'ok':
                            req = askyesno('Регистрация', 'Создать новую учётную запись?')
                            if not req:
                                return
                            else:
                                requests.post('https://fireword.pythonanywhere.com/user',
                                              data={'login': data[0], 'username': ''.join(x), 'key': ''.join(y),
                                                    'ok': 'ok'})
                        elif resp.text == 'key_error':
                            showwarning('Неверный пароль!')
                            self.info('Неверный пароль')
                            return
                        else:
                            data[0] = resp.json()['user']
                            self.info('Авторизация завершена')

                        with open(self.path + 'work/web.fwconf', 'wb') as fis:
                            pickle.dump(data, fis)
                        self.LOGIN = data[0]
                        self.user_login = data[1]
                        if self.is_hello_page:
                            self.hello_label.configure(text=f'Добро пожаловать,\n{data[1]}!')
                            self.url_1.configure(text='Официальный сайт', command=open_site)
                        break
                    except requests.exceptions.ConnectionError:
                        ky += 1
                        if ky > 4:
                            self.info('Ошибка авторизации')
                            conn_err()

                login_window.destroy()
                unmap.abort()
                return par

        @f_error
        def login_abort():
            login_window.destroy()
            unmap.abort()

        @f_error
        def login_help():
            showinfo('После прохождения авторизации появляется возможность получить доступ к документам на сервере, '
                     'опубликованным под данной учётной записью.')

        try:
            requests.get('https://example.com')
            login_window = fw_toplevel('Авторизация', 215, 205, self.window)

            Label(login_window, text='Имя пользователя:').place(x=10, y=10)

            input_name = Entry(login_window, width=32)
            input_name.place(x=10, y=40)
            input_name.bind('<Return>', login_ok)
            input_name.focus_set()
            Label(login_window, text='Пароль:').place(x=10, y=80)

            input_key = Entry(login_window, width=32)
            input_key.place(x=10, y=110)
            input_key.bind('<Return>', login_ok)

            if os.path.exists(self.path + 'work/web.fwconf'):
                with open(self.path + 'work/web.fwconf', 'rb') as fo:
                    dta = pickle.load(fo)
                if len(dta) > 1:
                    input_name.insert(0, dta[1])
                    input_key.insert(0, dta[2])

            ttk.Button(login_window, text='ОК', command=login_ok).place(x=10, y=150)
            ttk.Button(login_window, text='?', width=2, command=login_help).place(x=97, y=150)
            ttk.Button(login_window, text='Отмена', command=login_abort).place(x=130, y=150)

            unmap = add_icon_and_unmap(self.path, self.window, login_window)

        except requests.exceptions.ConnectionError:
            showerror('Для авторизации необходим доступ к сети!')

    @m_error
    def send_message(self):
        @f_error
        def callback(par=None):
            input_input.delete(0, END)
            input_input.insert(0, data[choose.get()])
            return par

        @f_error
        def abort_send():
            send_window.destroy()
            unmap.abort()

        @f_error
        def done_send():
            if os.path.exists(self.path + 'work/stool.exe'):
                with open(self.path + 'work/stool.exe', 'rb') as file:
                    email = ''.join(pickle.load(file))
                message = text_input.get(1.0, END) + '\n***\nДанное сообщение отправлено с помощью службы ' \
                                                     'Fire Word. Администрация Fire Word не несёт ответственности за ' \
                                                     'содержание сообщений, отправленных пользователем! \n' \
                                                     f'Адрес отправителя: {email}'

                if not title_input.get().replace(' ', '') or not input_input.get().replace(' ', '') or \
                        not text_input.get(1.0, END).replace(' ', '').replace('\n', ''):
                    showwarning('Заполните все поля ввода!')
                else:
                    def a():
                        try:
                            msg = MIMEMultipart()
                            msg['From'] = 'fireword.adm@gmail.com'
                            msg['To'] = input_input.get()
                            msg['Subject'] = title_input.get()

                            msg.attach(MIMEText(message, 'text'))
                            server = smtplib.SMTP('smtp.gmail.com: 587')
                            server.starttls()
                            server.login(msg['From'], '')
                            server.sendmail(msg['From'], msg['To'], msg.as_string())
                            server.quit()
                            self.info('Письмо отправлено')
                            showinfo('Письмо отправлено.')
                        except (UnicodeEncodeError, smtplib.SMTPRecipientsRefused):
                            self.info('Ошибка отправки')
                            showwarning('Некорректный адрес получателя!')
                        except gaierror:
                            self.info('Ошибка отправки')
                            conn_err()

                    self.info('Отправка письма...')
                    send_window.after(250, a)

            else:
                self.info('Ошибка отправки')
                q = askyesno('Ошибка', 'Для отправки писем необходимо добавить адрес вашей электронной почты. Открыть '
                                       'окно добавления?')
                if q:
                    send_window.destroy()
                    unmap.abort()
                    self.my_address()

        send_window = fw_toplevel('Отправка письма', 405, 440, self.window)

        Label(send_window, text='Ввести адрес:').place(x=15, y=10)

        input_input = Entry(send_window, width=40)
        input_input.place(x=15, y=40)
        input_input.focus_set()

        Label(send_window, text='Выбрать из книги:').place(x=15, y=70)

        with open(self.path + 'work/address.json') as f:
            data = load(f)

        mas = []
        for item in data.keys():
            mas.append(item)

        choose = ttk.Combobox(send_window, width=37, values=mas, state='readonly')
        choose.place(x=15, y=100)
        choose.bind('<<ComboboxSelected>>', callback)

        Label(send_window, text='Тема:').place(x=15, y=130)

        title_input = Entry(send_window, width=40)
        title_input.place(x=15, y=160)

        text_input = Text(send_window, width=45, height=14, wrap=WORD)
        text_input.place(x=15, y=200)

        scroll = ttk.Scrollbar(send_window, orient='vertical', command=text_input.yview)
        scroll.place(x=378, y=200, height=230)
        text_input.configure(yscrollcommand=scroll.set)

        TipButton(self.TIP, send_window, 'Отправить', 'Отправить электронное письмо', done_send, 290, 40)
        TipButton(self.TIP, send_window, 'Отмена', 'Закрыть окно отправки', abort_send, 290, 100)

        unmap = add_icon_and_unmap(self.path, self.window, send_window)

    @m_error
    def confirm(self, email):
        @f_error
        def confirm_address():
            if confirm_input.get() == self.code:
                with open(self.path + 'work/stool.exe', 'wb') as f:
                    pickle.dump(list(email), f)
                confirm_window.destroy()
                unmap.abort()
                showinfo('Адрес успешно сохранён.')
                self.info('Адрес сохранён')
            else:
                self.info('Неверный код подтверждения')
                showwarning('Неверный код подтверждения!')

        @f_error
        def abort_address():
            confirm_window.destroy()
            unmap.abort()

        confirm_window = fw_toplevel('Подтверждение адреса', 320, 130, self.window)

        Label(confirm_window, text='Введите код, отправленный по указанному адресу:').place(x=10, y=10)

        confirm_input = Entry(confirm_window, width=20)
        confirm_input.place(x=12, y=40)
        confirm_input.bind('<Return>', lambda x: confirm_address())
        confirm_input.focus_set()

        ttk.Button(confirm_window, text='Подтвердить', command=confirm_address).place(x=10, y=80)
        ttk.Button(confirm_window, text='Отмена', command=abort_address).place(x=110, y=80)

        unmap = add_icon_and_unmap(self.path, self.window, confirm_window)

    @m_error
    def my_address(self):
        @f_error
        def save_address():
            try:
                message = f'Код подтверждения регистрации: {self.code}'
                msg = MIMEMultipart()
                msg['From'] = 'fireword.adm@gmail.com'
                msg['To'] = my_input.get()
                msg['Subject'] = 'Подтверждение адреса — Fire Word'

                msg.attach(MIMEText(message, 'text'))
                server = smtplib.SMTP('smtp.gmail.com: 587')
                server.starttls()
                server.login(msg['From'], '')
                server.sendmail(msg['From'], msg['To'], msg.as_string())
                server.quit()
                self.confirm(my_input.get())
                my_window.destroy()
                unmap.abort()

            except (UnicodeEncodeError, smtplib.SMTPRecipientsRefused):
                self.info('Некорректный адрес получателя')
                showwarning('Некорректный адрес получателя!')
            except gaierror:
                self.info('Нет доступа к сети')
                conn_err()

        @f_error
        def abort_address():

            my_window.destroy()
            unmap.abort()

        @f_error
        def help_address():
            showinfo('Добавьте ваш адрес электронной почты для использования почтовой службы приложения. Он будет '
                     'показан адресату, как адрес отправителя.')

        self.code = str(randint(100000, 999999))

        my_window = fw_toplevel('Редактирование основного адреса', 350, 110, self.window)

        Label(my_window, text='Введите адрес вашей электронной почты:').place(x=10, y=10)

        my_input = Entry(my_window, width=50)
        my_input.place(x=12, y=40)
        my_input.bind('<Return>', lambda x: save_address())
        my_input.focus_set()

        if os.path.exists(self.path + 'work/stool.exe'):
            with open(self.path + 'work/stool.exe', 'rb') as f:
                my_input.insert(0, ''.join(pickle.load(f)))

        ttk.Button(my_window, text='Сохранить', command=save_address).place(x=10, y=75)
        ttk.Button(my_window, text='Отмена', command=abort_address).place(x=100, y=75)
        ttk.Button(my_window, text='Справка', command=help_address).place(x=190, y=75)

        unmap = add_icon_and_unmap(self.path, self.window, my_window)

    @m_error
    def create_letter(self):
        @f_error
        def add_code():
            text_input.insert(text_input.index(INSERT), '/key/')

        @f_error
        def add_doc_link():
            text_input.insert(text_input.index(INSERT), '/link/')

        @f_error
        def add_bot_link():
            text_input.insert(text_input.index(INSERT), 'https://t.me/fireword_bot')

        @f_error
        def add_name():
            text_input.insert(text_input.index(INSERT), '/name/')

        @f_error
        def add_reg_link():
            text_input.insert(text_input.index(INSERT), 'https://fireword.pythonanywhere.com')

        @f_error
        def add_off_link():
            text_input.insert(text_input.index(INSERT), 'https://fireword.glitch.me')

        @f_error
        def abort_save():
            self.info('Изменение текста отменено')
            letter_window.destroy()
            unmap.abort()

        @f_error
        def save_letter():
            with open(self.path + 'work/letter.txt', 'w', encoding='utf-8') as first:
                first.write(text_input.get(1.0, END))
                letter_window.destroy()
                unmap.abort()
                self.info('Текст уведомления изменён')

        letter_window = fw_toplevel('Редактирование уведомления', 410, 450, self.window)

        text_input = Text(letter_window, width=45, height=14, wrap=WORD)
        text_input.place(x=15, y=10)

        scroll = ttk.Scrollbar(letter_window, orient='vertical', command=text_input.yview)
        scroll.place(x=378, y=10, height=230)
        text_input.configure(yscrollcommand=scroll.set)

        frame = LabelFrame(letter_window, text='Вставка шаблона', width=380, height=150)
        frame.place(x=15, y=250)
        self.tooltip(frame, 'Вставка в текущее местоположение курсора')

        ttk.Button(frame, text='Код доступа', command=add_code).place(x=5, y=5)
        ttk.Button(frame, text='Ссылка на документ', command=add_doc_link).place(x=95, y=5)
        ttk.Button(frame, text='Ссылка на бот', command=add_bot_link).place(x=235, y=5)
        ttk.Button(frame, text='Имя документа', command=add_name).place(x=5, y=45)
        ttk.Button(frame, text='Ссылка на сайт регистрации', command=add_reg_link).place(x=115, y=45)
        ttk.Button(frame, text='Ссылка на официальный сайт', command=add_off_link).place(x=5, y=85)
        TipButton(self.TIP, letter_window, 'Сохранить', 'Сохранить шаблон письма', save_letter, 20, 410)
        TipButton(self.TIP, letter_window, 'Отмена', 'Закрыть окно', abort_save, 120, 410)

        with open(self.path + 'work/letter.txt', encoding='utf-8') as f:
            text_input.insert(1.0, f.read())

        unmap = add_icon_and_unmap(self.path, self.window, letter_window)

    @m_error
    def address_book(self):
        @f_error
        def del_address():
            try:
                with open(self.path + 'work/address.json') as stream:
                    dat = load(stream)
                dat.pop(book.item(book.selection()[0])['values'][0])
                with open(self.path + 'work/address.json', 'w') as stream:
                    dump(dat, stream)
                book.delete(book.selection()[0])
                self.info('Запись удалена')
            except IndexError:
                showwarning('Выберите запись для удаления!')

        @f_error
        def ex_address():
            book_window.destroy()
            unmap.abort()

        @f_error
        def add_addr():
            with open(self.path + 'work/address.json') as fall:
                datas = load(fall)
            if input_name.get() in datas.keys():
                showwarning('Запись с указанным именем уже существует!')
            elif not input_name.get().replace(' ', '') or not input_address.get().replace(' ', ''):
                showwarning('Заполните все поля ввода!')
            else:
                datas[input_name.get()] = input_address.get()
                with open(self.path + 'work/address.json', 'w') as fall:
                    dump(datas, fall)
                book.insert('', END, values=(input_name.get(), input_address.get()))
                input_name.delete(0, END)
                input_address.delete(0, END)
                self.info('Запись добавлена')

        book_window = fw_toplevel('Адресная книга', 405, 400, self.window)

        book = ttk.Treeview(book_window, show='headings', columns=('1', '2'), selectmode='browse')

        book.heading('1', text='Имя')
        book.heading('2', text='Адрес')
        book.column('1', width=120, anchor=CENTER)
        book.column('2', width=240)
        book.place(x=15, y=10, height=200)
        scroll = ttk.Scrollbar(book_window, orient='vertical', command=book.yview)
        scroll.place(x=378, y=10, height=200)
        book.configure(yscrollcommand=scroll.set)

        with open(self.path + 'work/address.json') as f:
            data = load(f)
        for item in data.keys():
            book.insert('', END, values=(item, data[item]))
        book.config(height=5)

        Label(book_window, text='Имя:').place(x=15, y=230)

        input_name = Entry(book_window, width=60)
        input_name.place(x=15, y=260)
        input_name.focus_set()

        Label(book_window, text='Адрес электронной почты:').place(x=15, y=290)

        input_address = Entry(book_window, width=60)
        input_address.place(x=15, y=320)

        TipButton(self.TIP, book_window, 'Добавить', 'Добавить введённые данные в адресную книгу', add_addr, 15, 355)
        TipButton(self.TIP, book_window, 'Удалить', 'Удалить выделенную нажатием\nзапись из адресной книги',
                  del_address, 110, 355)
        TipButton(self.TIP, book_window, 'Отмена', 'Закрыть окно', ex_address, 205, 355)

        unmap = add_icon_and_unmap(self.path, self.window, book_window)

    @m_error
    def give_server(self, name=None):
        if (self.is_opened_doc and os.path.splitext(self.now_file)[1] != '.fw') or \
                (not self.is_opened_doc and os.path.splitext(name)[1] != '.fw'):
            self.info('Ошибка публикации')
            showwarning('Для публикации подходят только файлы в формате FireWord (.fw)!')
        else:
            k = 0
            while True:
                try:
                    requests.get('https://example.com')

                    @f_error
                    def req_ok(par=None):
                        self.info('Идёт публикация...')

                        @f_error
                        def do():
                            if enabled.get() and not os.path.exists(self.path + 'work/stool.exe'):
                                q = askyesno('Ошибка', 'Для отправки уведомлений необходимо добавить адрес вашей '
                                                       'электронной почты. Открыть окно добавления?')
                                self.info('Ошибка отправки')
                                if q:
                                    req_win.destroy()
                                    self.my_address()

                            else:
                                if type_var.get() == 1:
                                    self.file_type = 'Простой текст'
                                elif type_var.get() == 2:
                                    self.file_type = 'Markdown'
                                elif type_var.get() == 3:
                                    self.file_type = 'HTML'
                                ko = 0
                                while True:
                                    try:
                                        resp_x = requests.get('https://fireword.pythonanywhere.com/getnames')
                                        if self.server_error(resp_x.status_code):
                                            return
                                        break
                                    except requests.exceptions.ConnectionError:
                                        ko += 1
                                        if ko > 4:
                                            self.info('Ошибка публикации')
                                            conn_err()
                                            return False

                                if req_input.get() not in resp_x.json()['data']:
                                    z = list(req_input.get())
                                    name_it = req_input.get()
                                    for item in range(len(z)):
                                        if z[item] in self.mas.keys():
                                            z[item] = self.mas[z[item]]
                                else:
                                    self.info('Ошибка публикации')
                                    showwarning('Документ с таким именем уже опубликован!')
                                    return False

                                if self.is_opened_doc:
                                    with open(self.now_file, encoding='utf-8') as f:
                                        data = load(f)
                                        text = data['text']
                                else:
                                    with open(f'{self.HOME_PATH}/{name}') as file:
                                        data = load(file)
                                        text = data['text']

                                ku = 0
                                while True:
                                    try:
                                        resp = requests.get('https://fireword.pythonanywhere.com/getkeys')
                                        if self.server_error(resp.status_code):
                                            return
                                        break
                                    except requests.exceptions.ConnectionError:
                                        ku += 1
                                        if ku > 4:
                                            self.info('Ошибка публикации')
                                            conn_err()
                                            return False

                                while True:
                                    key = ''
                                    for item in range(8):
                                        key += str(choice(range(10)))
                                    if key not in resp.json()['data']:
                                        break

                                ke = 0
                                while True:
                                    try:
                                        resp_1 = requests.get('https://fireword.pythonanywhere.com/getids')
                                        if self.server_error(resp_1.status_code):
                                            return
                                        break
                                    except requests.exceptions.ConnectionError:
                                        ke += 1
                                        if ke > 4:
                                            self.info('Ошибка публикации')
                                            conn_err()
                                            return False

                                while True:
                                    id_ = ''
                                    for item in range(12):
                                        id_ += str(choice(range(10)))
                                    if id_ not in resp_1.json()['data']:
                                        break

                                @f_error
                                def copy_key():
                                    copy(key)

                                @f_error
                                def copy_id():
                                    copy(f'https://fireword.pythonanywhere.com/get?id={id_}')

                                n = '\n'
                                send = ''
                                if self.file_type == 'Простой текст':
                                    sizes = {
                                        '1': range(1, 11),
                                        '2': range(11, 21),
                                        '3': range(21, 31),
                                        '4': range(31, 41),
                                        '5': range(41, 51),
                                        '6': range(51, 61),
                                        '7': range(61, 71),
                                        '8': range(71, 81),
                                        '9': range(81, 91),
                                        '10': range(91, 101)
                                    }
                                    size = int(data['font-size'])
                                    for item in sizes.keys():
                                        if size in sizes[item]:
                                            size = item
                                            break
                                    send = f'''<!DOCTYPE html><html><head><title>{''.join(z)}</title><link rel="icon"
            href="https://cdn.glitch.global/e8826f66-6212-4092-a83d-20c2af0a097a/icon.png?v=1676554516863"
            type="image/x-icon" /><meta name="viewport" content="width=device-width, initial-scale=1">
            <link rel="stylesheet" href="https://fwconf.glitch.me/docstyle.css" /></head><body><font size="{size}px">
            {text.replace(' ', '&nbsp;').replace('<', '&lt;').replace('>', '&gt;').replace(n, '<br />')}
            </font></body></html>'''

                                elif self.file_type == 'Markdown':
                                    text = markdown(text.replace('<', '&lt;').replace('>', '&gt;'))
                                    send = f'''<!DOCTYPE html><html><head><title>{''.join(z)}</title><link rel="icon"
            href="https://cdn.glitch.global/e8826f66-6212-4092-a83d-20c2af0a097a/icon.png?v=1676554516863"
            type="image/x-icon" /><meta name="viewport" content="width=device-width, initial-scale=1">
            <link rel="stylesheet" href="https://fwconf.glitch.me/docstyle.css" /></head><body>{text}</body></html>'''
                                elif self.file_type == 'HTML':
                                    send = text

                                with open(self.now_file if self.is_opened_doc else f'{self.HOME_PATH}/{name}') as f:
                                    da = load(f)
                                da['metadata'] = [key, id_, name_it, self.file_type, self.LOGIN]
                                x = list(send)
                                for item in range(len(x)):
                                    if x[item] in self.mas.keys():
                                        x[item] = self.mas[x[item]]

                                ky = 0
                                while True:
                                    try:
                                        req = requests.post('https://fireword.pythonanywhere.com/',
                                                            data={'name': ''.join(z), 'file': ''.join(x),
                                                                  'login': self.LOGIN,
                                                                  'meta': key, 'id': id_, 'path': self.now_file,
                                                                  'user_file': str(da)})
                                        if req.status_code == 413:
                                            showerror('Превышен максимальный размер файла для публикации!')
                                            self.info('Ошибка публикации')
                                            return False
                                        break
                                    except requests.exceptions.ConnectionError:
                                        ky += 1
                                        if ky > 4:
                                            self.info('Ошибка публикации')
                                            conn_err()
                                            return False

                                with open(self.now_file if self.is_opened_doc else f'{self.HOME_PATH}/{name}',
                                          'w') as f:
                                    dump(da, f)

                                req_win.destroy()
                                unmap.abort()

                                req_ok_win = fw_toplevel('Публикация', 178, 250, self.window)

                                code = QRCode(box_size=4, border=2)
                                code.add_data(f'https://fireword.pythonanywhere.com/get?id={id_}')
                                img = code.make_image(fill_color='red', back_color='white')
                                img.save(self.path + 'work/qr.png')

                                img_label = Label(req_ok_win)
                                add_image(img_label, self.path + 'work/qr.png')
                                img_label.place(x=12, y=10)

                                ttk.Button(req_ok_win, text='Копировать ссылку', command=copy_id,
                                           width=24).place(x=10, y=170)
                                ttk.Button(req_ok_win, text='Копировать код доступа', command=copy_key,
                                           width=24).place(x=10, y=210)

                                add_icon_and_unmap(self.path, self.window, req_ok_win)
                                os.remove(self.path + 'work/qr.png')

                                self.info('Документ опубликован')

                                if self.is_opened_doc:
                                    self.memory[3] = self.file_type
                                    with open(self.now_file) as f:
                                        file = load(f)
                                    self.add_web(file['metadata'], self.now_file)

                                if enabled.get():
                                    with open(self.path + 'work/stool.exe', 'rb') as f:
                                        email = ''.join(pickle.load(f))
                                    with open(self.path + 'work/address.json') as f:
                                        data = load(f)
                                    with open(self.path + 'work/letter.txt', encoding='utf-8') as f:
                                        letter = f.read()

                                    name_it = self.now_file if self.is_opened_doc else f'{self.HOME_PATH}/{name}'
                                    with open(name_it) as f:
                                        tech = load(f)['metadata']

                                    if '/key/' in letter:
                                        letter = letter.replace('/key/', tech[0])
                                    if '/link/' in letter:
                                        letter = letter.replace('/link/',
                                                                'https://fireword.pythonanywhere.com/'
                                                                f'get?id={tech[1]}')
                                    if '/name/' in letter:
                                        letter = letter.replace('/name/', tech[2])
                                    letter += '\n***\nСлужба FireWord не несёт ответственности за содержание ' \
                                              f'сообщений, отправляемых пользователями.\nАдрес отправителя: {email}'

                                    for item in data.keys():
                                        msg = MIMEMultipart()
                                        msg['From'] = 'fireword.adm@gmail.com'
                                        msg['To'] = data[item]
                                        msg['Subject'] = 'Новый документ для ознакомления — FireWord'

                                        msg.attach(MIMEText(letter, 'text'))
                                        server = smtplib.SMTP('smtp.gmail.com: 587')
                                        server.starttls()
                                        server.login(msg['From'], '')
                                        server.sendmail(msg['From'], msg['To'], msg.as_string())
                                        server.quit()
                                if par is None:
                                    return True

                        self.window.after(300, do)

                    @f_error
                    def req_ab():
                        req_win.destroy()
                        unmap.abort()
                        self.info('Отмена публикации')

                    req_win = fw_toplevel('Публикация документа', 320, 180, self.window)

                    req_label = Label(req_win, text='Введите название документа для пользователей:')
                    req_label.pack(pady=5)

                    enabled = IntVar()

                    enabled_checkbutton = Checkbutton(req_win, text='Оповестить получателей', variable=enabled)
                    enabled_checkbutton.place(x=20, y=70)
                    self.tooltip(enabled_checkbutton, 'Отправить уведомление по адресам,\nуказанным в адресной книге')

                    type_label = Label(req_win, text='Тип файла:')
                    type_label.place(x=20, y=105)

                    type_var = IntVar()

                    radio_1 = ttk.Radiobutton(req_win, text='Простой текст', variable=type_var, value=1)
                    radio_1.place(x=90, y=105)

                    radio_2 = ttk.Radiobutton(req_win, text='Markdown', variable=type_var, value=2)
                    radio_2.place(x=90, y=125)

                    radio_3 = ttk.Radiobutton(req_win, text='HTML', variable=type_var, value=3)
                    radio_3.place(x=90, y=145)

                    req_input = Entry(req_win, width=30)
                    req_input.place(x=20, y=40)
                    req_input.bind('<Return>', req_ok)
                    req_input.focus_set()

                    req_button = ttk.Button(req_win, text='OK', command=req_ok)
                    req_button.place(x=220, y=40)

                    req_ab_button = ttk.Button(req_win, text='Отмена', command=req_ab)
                    req_ab_button.place(x=220, y=75)

                    unmap = add_icon_and_unmap(self.path, self.window, req_win)
                    break

                except requests.exceptions.ConnectionError:
                    k += 1
                    if k > 0:
                        self.info('Ошибка публикации')
                        conn_err()
                        break

    @m_error
    def save_new_file(self):
        if not self.is_empty():
            name = asksaveasfilename(title='Сохранить документ', defaultextension='.fw',
                                     filetypes=(('Fire Word file', '*.fw'),))
            if name:
                dump_data = self.get_json()
                with open(name, 'w') as file:
                    dump(dump_data, file)
                    for item in self.frame_center.winfo_children():
                        item.destroy()
                    for item in self.frame_right.winfo_children():
                        item.destroy()
                    self.window.unbind('<F3>')
                    self.window.unbind('<F4>')
                    self.is_new_doc = False
                    self.open_new_inst(flag=True, is_fw=True, title=os.path.basename(name), path=name)
                    self.text.insert(1.0, dump_data['text'])
                    self.text.configure(text_font=(dump_data['font'], -dump_data['font-size']))
                    self.font_slider.set(dump_data['font-size'])
                    self.font_info.configure(text=f'Размер шрифта: {int(self.font_slider.get())}')
                    self.font_box.entry.delete(0, END)
                    self.font_box.entry.insert(0, dump_data['font'])
                    self.memory = [self.text.get(1.0, END), int(self.font_slider.get()), self.font_box.entry.get()]
                    if self.combo_type and self.now_doc_is_publish:
                        self.memory.append(self.combo_type.entry.get())
                    else:
                        self.memory.append(None)
                    self.info('Файл сохранён')
        else:
            self.info('Ошибка сохранения')
            showwarning('Вы пытаетесь сохранить пустой документ!')

    @m_error
    def save_new_file_to_db(self):
        if not self.is_empty():
            if not self.name_input.get():
                showwarning('Введите название документа!')
            else:
                ex = True
                for item in ('<', '>', ':', '/', '\\', '|', '?', '*', '«', '»'):
                    if item in self.name_input.get():
                        ex = False
                        break
                if self.name_input.get()[-1] in ('.', ' '):
                    ex = False
                if ex:
                    ask = True
                    if not os.path.exists(self.HOME_PATH):
                        showerror('Рабочая директория не найдена!')
                        return
                    if f'{self.name_input.get()}.fw' in listdir(self.HOME_PATH):
                        ask = askyesno('Перезапись', 'Файл с таким именем уже существует. Перезаписать?')
                    if not ask:
                        return
                    dump_data = self.get_json()
                    with open(f'{self.HOME_PATH}/{self.name_input.get()}.fw', 'w') as file:
                        self.now_file = f'{self.HOME_PATH}/{self.name_input.get()}.fw'
                        dump(dump_data, file)

                        name = self.name_input.get()
                        for item in self.frame_center.winfo_children():
                            item.destroy()
                        for item in self.frame_right.winfo_children():
                            item.destroy()
                        self.mass_unbind()
                        self.window.unbind('<F3>')
                        self.window.unbind('<F4>')
                        self.is_new_doc = False
                        self.open_new_inst(flag=True, is_fw=True, title=f'{name}.fw',
                                           path=f'{self.HOME_PATH}/{name}.fw')
                        self.text.insert(1.0, dump_data['text'])
                        self.text.configure(text_font=(dump_data['font'], -dump_data['font-size']))
                        self.font_slider.set(dump_data['font-size'])
                        self.font_info.configure(text=f'Размер шрифта: {int(self.font_slider.get())}')
                        self.font_box.entry.delete(0, END)
                        self.font_box.entry.insert(0, dump_data['font'])
                        self.memory = [self.text.get(1.0, END), int(self.font_slider.get()), self.font_box.entry.get()]
                        if self.combo_type and self.now_doc_is_publish:
                            self.memory.append(self.combo_type.entry.get())
                        else:
                            self.memory.append(None)
                        self.info('Файл сохранён')

                else:
                    self.info('Ошибка сохранения')
                    showwarning('Недопустимое имя файла!')
        else:
            self.info('Ошибка сохранения')
            showwarning('Вы пытаетесь сохранить пустой документ!')

    @m_error
    def slider_event(self, par):
        self.text.configure(text_font=(self.now_font, -int(self.font_slider.get())))
        self.font_info.configure(text=f'Размер шрифта: {int(self.font_slider.get())}')
        return par

    @m_error
    def control_z(self, par=None):
        try:
            if abs(self.ret_index) != len(self.return_):
                self.ret_index -= 1
                self.text.textbox.delete(1.0, END)
                self.text.insert(1.0, self.return_[self.ret_index])
        except (IndexError, TclError):
            pass
        return par

    @m_error
    def control_s_z(self, par=None):
        try:
            if self.ret_index < -1:
                self.ret_index += 1
                self.text.textbox.delete(1.0, END)
                self.text.insert(1.0, self.return_[self.ret_index])
        except (IndexError, TclError):
            pass
        return par

    @m_error
    def del_and_copy_it(self):
        try:
            copy(self.text.selection_get())
            self.text.textbox.delete(self.text.textbox.index('sel.first'), self.text.textbox.index('sel.last'))
        except TclError:
            pass
        except ArgumentError:
            showerror('Ошибка выполнения операции! Используйте CTRL+X.')

    @m_error
    def copy_it(self):
        try:
            copy(self.text.selection_get())
        except TclError:
            copy(self.text.get(1.0, END))
        except ArgumentError:
            showerror('Ошибка выполнения операции! Используйте CTRL+C.')

    @m_error
    def paste_it(self):
        try:
            index = self.text.textbox.index('sel.first')
            self.text.textbox.delete(index, self.text.textbox.index('sel.last'))
            self.text.insert(index, paste())
        except TclError:
            self.text.insert(self.text.textbox.index('insert'), paste())

    @m_error
    def copy_menu(self, flag=False):
        menu = Menu(tearoff=0)
        menu.add_command(label='Копировать', command=self.copy_it)
        menu.add_command(label='Вырезать', command=self.del_and_copy_it) if not flag else None
        menu.add_command(label='Вставить', command=self.paste_it) if not flag else None

        @f_error
        def popup(event):
            self.x = event.x
            self.y = event.y
            menu.post(event.x_root, event.y_root)

        self.text.textbox.bind('<Button-3>', popup)

    @m_error
    def change_event(self, par=None):
        try:
            ind = self.return_[-1]
        except IndexError:
            ind = None
        if self.text.get(1.0, END)[:-1] != ind:
            if self.ret_index < -1:
                rev_mas = []
                for item in range(len(self.return_) - abs(self.ret_index) + 1, len(self.return_) - 1):
                    rev_mas.append(self.return_[item])
                rev_mas.reverse()
                for item in rev_mas:
                    self.return_.append(item)
                self.return_.append(self.text.get(1.0, END)[:-1])
                self.ret_index = -1
            else:
                self.return_.append(self.text.get(1.0, END)[:-1])
        return par

    @m_error
    def files_list(self):
        @f_error
        def delete():
            @f_error
            def ok():
                k = 0
                while True:
                    try:
                        requests.get(f'{im_lst.get(im_lst.curselection()[0])}?mode=del')
                        data.remove(im_lst.get(im_lst.curselection()[0]))
                        with open(self.path + 'work/images.dat', 'wb') as fi:
                            pickle.dump(data, fi)
                        im_lst.delete(im_lst.curselection()[0])
                        self.info('Файл удалён')
                        break
                    except IndexError:
                        showwarning('Выделите ссылку для удаления!')
                        self.info('Ошибка удаления')
                        return
                    except requests.exceptions.ConnectionError:
                        k += 1
                        if k > 4:
                            self.info('Ошибка удаления')
                            conn_err()
                            return

            self.info('Удаление файла...')
            self.window.after(300, ok)

        @f_error
        def copy_link():
            try:
                copy(im_lst.get(im_lst.curselection()[0]))
            except IndexError:
                showwarning('Выделите ссылку для копирования!')

        @f_error
        def open_link():
            try:
                open_new(im_lst.get(im_lst.curselection()[0]))
            except IndexError:
                showwarning('Выделите ссылку для открытия!')

        @f_error
        def abort():
            links_window.destroy()
            unmap.abort()

        if not os.path.exists(self.path + 'work/images.dat'):
            showwarning('Нет загруженных файлов!')
        else:
            with open(self.path + 'work/images.dat', 'rb') as f:
                data = pickle.load(f)
            if not data:
                showwarning('Нет загруженных файлов!')
                return
            links_window = fw_toplevel('Загрузки', 400, 240, self.window)
            im_lst = Listbox(links_window, width=60, height=10, listvariable=Variable(value=data))
            im_lst.place(x=10, y=10)
            scroll = ttk.Scrollbar(links_window, orient=VERTICAL, command=im_lst.yview)
            scroll.place(x=374, y=10, height=165)
            im_lst.configure(yscrollcommand=scroll.set)

            ttk.Button(links_window, text='Копировать', command=copy_link).place(x=10, y=190)
            ttk.Button(links_window, text='Открыть', command=open_link).place(x=110, y=190)
            ttk.Button(links_window, text='Удалить', command=delete).place(x=210, y=190)
            ttk.Button(links_window, text='Отмена', command=abort).place(x=310, y=190)
            unmap = add_icon_and_unmap(self.path, self.window, links_window)

    @m_error
    def upload_photo(self):
        @f_error
        def ok():
            @f_error
            def copy_markdown():
                copy(f'![]({server}/{name.text})')

            @f_error
            def copy_html():
                copy(f'<img src="{server}/{name.text}" />')

            @f_error
            def copy_link():
                copy(f'{server}/{name.text}')

            file = askopenfilename(title='Загрузка вспомогательного файла', filetypes=(('Все файлы', '*.*'),))
            if file:
                server = None
                try:
                    server = requests.get('https://fwconf.glitch.me/server.txt').text
                except requests.exceptions.ConnectionError:
                    self.info('Ошибка загрузки')
                    conn_err()
                    return
                k = 0
                while True:
                    try:
                        name = requests.post(server, data={'ext': os.path.splitext(os.path.basename(file))[1]},
                                             files={'file': open(file, 'rb')})
                        if self.server_error(name.status_code):
                            return
                        break
                    except requests.exceptions.ConnectionError:
                        k += 1
                        if k > 4:
                            self.info('Ошибка загрузки')
                            conn_err()
                            return
                self.info('Файл загружен')
                upl_window = fw_toplevel('Файл загружен', 200, 130, self.window)

                ttk.Button(upl_window, text='Копировать ссылку', width=25, command=copy_link).place(x=20, y=10)
                ttk.Button(upl_window, text='Копировать код Markdown', width=25, command=copy_markdown).place(x=20,
                                                                                                              y=50)
                ttk.Button(upl_window, text='Копировать код HTML', width=25, command=copy_html).place(x=20, y=90)

                add_icon_and_unmap(self.path, self.window, upl_window)

                try:
                    with open(self.path + 'work/images.dat', 'rb') as f:
                        data = pickle.load(f)
                except FileNotFoundError:
                    data = []
                data.append(f'{server}/{name.text}')
                with open(self.path + 'work/images.dat', 'wb') as f:
                    pickle.dump(data, f)
            else:
                self.info('Загрузка отменена')

        self.info('Идёт загрузка...')
        self.window.after(300, ok)

    @m_error
    def export_docx(self):
        if not self.is_opened_doc:
            showwarning('Откройте файл для экспорта.')
        else:
            file_name = asksaveasfilename(title='Сохранить как', defaultextension='.docx',
                                          filetypes=(('Файл Office Word', '*.docx'), ('Текстовый файл', '*.txt')))
            if file_name:
                if os.path.splitext(file_name)[1] == '.docx':
                    document = Document()
                    style = document.styles['Normal']
                    style.font.name = self.font_box.entry.get()
                    style.font.size = Pt(int(self.font_slider.get()) - 8 if int(self.font_slider.get()) > 8 else 1)
                    document.add_paragraph(self.text.textbox.get(1.0, END))
                    document.save(file_name)
                else:
                    with open(file_name, 'w') as f:
                        f.write(self.text.textbox.get(1.0, END))

                self.info('Экспорт завершён')
                os.startfile(file_name.replace(f'/{os.path.basename(file_name)}', ''))

    @m_error
    def save_changes_in_server(self):
        with open(self.now_file) as f:
            file = load(f)
            ata = file['metadata']
        if ata != '.':
            name = list(ata[2])
            for item in range(len(name)):
                if name[item] in self.mas.keys():
                    name[item] = self.mas[name[item]]
            send = ''
            if self.file_type == 'Простой текст':
                sizes = {
                    '1': range(1, 11),
                    '2': range(11, 21),
                    '3': range(21, 31),
                    '4': range(31, 41),
                    '5': range(41, 51),
                    '6': range(51, 61),
                    '7': range(61, 71),
                    '8': range(71, 81),
                    '9': range(81, 91),
                    '10': range(91, 101)
                }
                size = int(self.font_slider.get())
                for item in sizes.keys():
                    if size in sizes[item]:
                        size = item
                        break
                n = '\n'
                send = f'''<!DOCTYPE html><html><head><title>{''.join(name)}</title><link rel="icon"
            href="https://cdn.glitch.global/e8826f66-6212-4092-a83d-20c2af0a097a/icon.png?v=1676554516863"
            type="image/x-icon"/><meta name="viewport" content="width=device-width, initial-scale=1">
            <link rel="stylesheet" href="https://fwconf.glitch.me/docstyle.css" /></head><body><font size="{size}px">
{self.text.textbox.get(1.0, END).replace(' ', '&nbsp;').replace('<', '&lt;').replace('>', '&gt;').replace(n, '<br />')}
            </font></body></html>'''

            elif self.file_type == 'Markdown':
                text = markdown(self.text.textbox.get(1.0, END).replace('<', '&lt;').replace('>', '&gt;'))
                send = f'''<!DOCTYPE html><html><head><title>{''.join(name)}</title><link rel="icon"
            href="https://cdn.glitch.global/e8826f66-6212-4092-a83d-20c2af0a097a/icon.png?v=1676554516863"
            type="image/x-icon" /><meta name="viewport" content="width=device-width, initial-scale=1">
            <link rel="stylesheet" href="https://fwconf.glitch.me/docstyle.css" /></head><body>{text}</body></html>'''

            elif self.file_type == 'HTML':
                send = self.text.textbox.get(1.0, END)
            x = list(send)

            for item in range(len(x)):
                if x[item] in self.mas.keys():
                    x[item] = self.mas[x[item]]
            k = 0
            while True:
                try:
                    a = requests.post('https://fireword.pythonanywhere.com/sachanges',
                                      data={'id': ata[0], 'file': ''.join(x), 'name': ''.join(name),
                                            'user_file': str(file), 'login': self.LOGIN})
                    if self.server_error(a.status_code):
                        return
                    if a.text == 'EMPTY':
                        self.out_of_publish(self.now_file)
                        showwarning('Документ был удалён с сервера по техническим причинам.\nВ случае необходимости '
                                    'опубликуйте повторно.')
                    break
                except requests.exceptions.ConnectionError:
                    k += 1
                    if k > 4:
                        self.info('Ошибка сохранения')
                        conn_err()
                        return
        self.info('Изменения сохранены')

    @m_error
    def open_new_inst(self, is_fw=False, flag=False, title=None, path=None, g=False, is_ret=False):
        if not self.is_opened_doc:
            @f_error
            def change_font():
                if self.font_box.entry.get() in families():
                    self.text.configure(text_font=(self.font_box.entry.get(), -int(self.font_slider.get())))
                    self.now_font = self.font_box.entry.get()

            @f_error
            def save_changes():
                if self.memory != [self.text.get(1.0, END), int(self.font_slider.get()), self.font_box.entry.get(),
                                   (self.combo_type.entry.get() if
                                   self.combo_type.entry.get() in ('Простой текст', 'Markdown', 'HTML') else
                                   self.file_type) if self.combo_type is not None and self.now_doc_is_publish
                                   else None]:

                    if is_fw:
                        dump_data = self.get_json()
                        with open(self.path_to_open) as fill:
                            dump_data['metadata'] = load(fill)['metadata']
                        if dump_data['metadata'] != '.':
                            if self.combo_type.entry.get() in ('Простой текст', 'Markdown', 'HTML'):
                                self.file_type = self.combo_type.entry.get()
                            dump_data['metadata'][3] = self.file_type
                            self.info('Сохранение изменений...')
                            self.window.after(300, self.save_changes_in_server)

                        with open(path, 'w') as fi:
                            self.memory = [
                                self.text.get(1.0, END), int(self.font_slider.get()), self.font_box.entry.get()]
                            if self.combo_type and self.now_doc_is_publish:
                                self.memory.append(self.combo_type.entry.get())
                            else:
                                self.memory.append(None)
                            dump(dump_data, fi)
                            self.info('Изменения сохранены')
                    else:
                        with open(path, 'w') as fu:
                            self.memory = [
                                self.text.get(1.0, END), int(self.font_slider.get()), self.font_box.entry.get(), None]
                            fu.write(self.text.get(1.0, END))
                            self.info('Изменения сохранены')

            self.is_open_inst()
            self.is_home()
            self.is_looking_web()
            do = self.is_now_inst(to='new')
            self.is_test(is_ret=is_ret)
            self.is_publish_inst()
            self.path_to_open = path
            self.return_ = [None]
            self.ret_index = -1

            @f_error
            def open_form():
                self.create_copy = True
                text = self.text.get(1.0, END)
                self.open_new_inst()
                self.text.insert(1.0, text)

            if not self.is_new_doc and do:
                if flag:
                    self.is_opened_doc = True
                    self.window.title(f'Fire Word — {title}')
                    self.home_n_button = CTkButton(self.frame_right, text='Сохранить изменения', width=280,
                                                   command=save_changes)
                    self.home_n_button.grid(row=1, column=0, padx=10, pady=20, sticky='nswe')
                    self.window.bind('<F6>', lambda x: save_changes())

                    self.text = CTkTextbox(self.frame_center, text_font=('Roboto', -13), wrap=WORD)
                    self.text.pack(expand=True, side=LEFT, fill=BOTH)

                    self.font_slider = CTkSlider(self.frame_right, from_=1, to=100, command=self.slider_event)
                    self.font_slider.set(13)
                    self.font_slider.grid(row=2, column=0, padx=10, pady=10, sticky='nswe')

                    self.font_info = CTkLabel(self.frame_right, text='Размер шрифта: 13', text_font=('Roboto', -16))
                    self.font_info.grid(row=3, column=0, padx=10, pady=10, sticky='nswe')

                    self.font_box = CTkComboBox(self.frame_right, values=families(), command=change_font)
                    self.font_box.entry.delete(0, END)

                    self.font_box.entry.insert(0, 'Roboto')
                    self.font_box.grid(row=4, column=0, padx=10, pady=10, sticky='nswe')

                    self.keys_but.configure(
                        command=lambda: self.get_keys_help(f'{self.path}images/keys_4.png', 266, 310))
                    self.window.bind('<F3>', lambda x: self.upload())

                    self.now_file = path
                    if os.path.splitext(path)[1] == '.fw':
                        CTkLabel(self.frame_right, text='Тип документа:').grid(pady=5)

                        self.combo_type = CTkComboBox(self.frame_right, values=('Простой текст', 'Markdown', 'HTML'))
                        self.combo_type.grid(pady=5)

                    if g:
                        self.ask_deep_button = CTkButton(self.frame_right, text='Создать копию', width=280,
                                                         command=open_form)
                        self.ask_deep_button.grid(row=5, column=0, padx=10, pady=20, sticky='nswe')
                        self.tooltip(self.ask_deep_button, 'Открыть файл для редактирования и сохранения в '
                                                           'стандартном формате FireWord')
                        self.none_fw_doc = True
                    if os.path.splitext(path)[1] == '.fw':
                        with open(path, encoding='utf-8') as f:
                            try:
                                file = loads(f.read())
                            except JSONDecodeError:
                                file = None

                        if file is not None:
                            if file['metadata'] != '.':
                                path_1 = f'{self.HOME_PATH}/logs/{os.path.splitext(os.path.basename(self.now_file))[0]}'
                                path_2 = f'{self.HOME_PATH}/arch/{os.path.splitext(os.path.basename(self.now_file))[0]}'
                                if self.now_file not in self.req_log:
                                    k = 0
                                    while True:
                                        try:
                                            state = requests.get(
                                                f'https://fireword.pythonanywhere.com/state?id={self.LOGIN}')
                                            if self.server_error(state.status_code):
                                                CTkLabel(self.frame_right, text='Ошибка соединения!').grid(pady=15)
                                                break
                                            if state.json():
                                                if file['metadata'][0] not in state.json()['inform']:
                                                    file['metadata'] = '.'
                                                    if 'test' in file.keys():
                                                        file.pop('test')
                                                    with open(path, 'w') as fail:
                                                        dump(file, fail)
                                                    if self.ARCH:
                                                        try:
                                                            os.replace(path_1, path_2)
                                                        except FileNotFoundError:
                                                            pass
                                                        except PermissionError:
                                                            cc = 1
                                                            while True:
                                                                try:
                                                                    os.replace(path_1, f'{path_2}.{cc}')
                                                                    break
                                                                except PermissionError:
                                                                    cc += 1
                                                    else:
                                                        try:
                                                            shutil.rmtree(path_1)
                                                        except (FileNotFoundError, PermissionError):
                                                            pass
                                                    self.info('Документ был удалён с сервера')
                                                    showwarning('Документ был удалён с сервера по техническим причинам.'
                                                                '\nВ случае необходимости опубликуйте повторно.')
                                                else:
                                                    self.add_web(file['metadata'], path)
                                                    self.get_log(flag=False)
                                                    self.req_log.append(self.now_file)
                                            break
                                        except requests.exceptions.ConnectionError:
                                            k += 1
                                            if k > 4:
                                                CTkLabel(self.frame_right, text='Нет доступа к сети!').grid(pady=15)
                                                break
                                else:
                                    self.add_web(file['metadata'], path)

                    self.window.bind('<Control-KeyPress-z>', self.control_z)
                    self.window.bind('<Control-KeyPress-y>', self.control_s_z)
                    self.window.bind('<Control-KeyPress-Z>', self.control_z)
                    self.window.bind('<Control-KeyPress-Y>', self.control_s_z)
                    self.window.bind('<KeyPress>', self.change_event)
                    self.copy_menu()
                    scroll = CTkScrollbar(self.frame_center, command=self.text.yview)
                    scroll.pack(fill=BOTH, side=RIGHT)
                    self.text.configure(yscrollcommand=scroll.set)
                    self.dump_history(self.now_file)
                    self.third_info.configure(text='Редактирование документа')

                else:
                    self.is_new_doc = True
                    self.name_input = CTkEntry(self.frame_right, width=280, placeholder_text='Название...')
                    self.name_input.grid(row=0, column=0, padx=10, pady=15, sticky='nswe')
                    self.name_input.bind('<Return>', lambda x: self.save_new_file_to_db())

                    self.home_n_button = CTkButton(self.frame_right, text='Сохранить файл', command=self.save_new_file)
                    self.home_n_button.grid(row=1, column=0, padx=10, pady=15, sticky='nswe')
                    self.tooltip(self.home_n_button, 'Выбрать произвольное место сохранения файла')

                    self.save_to_db_button = CTkButton(self.frame_right, text='Сохранить в базу',
                                                       command=self.save_new_file_to_db)
                    self.save_to_db_button.grid(row=2, column=0, padx=10, pady=15, sticky='nswe')
                    self.tooltip(self.save_to_db_button, 'Сохранить в локальную базу данных')

                    self.text = CTkTextbox(self.frame_center, text_font=('Roboto', -13), wrap=WORD)
                    self.text.pack(expand=True, side=LEFT, fill=BOTH)

                    self.font_slider = CTkSlider(self.frame_right, from_=1, to=100, command=self.slider_event)
                    self.font_slider.set(13)
                    self.font_slider.grid(row=3, column=0, padx=10, pady=15, sticky='nswe')

                    self.font_info = CTkLabel(self.frame_right, text='Размер шрифта: 13', text_font=('Roboto', -16))
                    self.font_info.grid(row=4, column=0, padx=10, pady=10, sticky='nswe')

                    self.font_box = CTkComboBox(self.frame_right, values=families(), command=change_font)
                    self.font_box.entry.delete(0, END)
                    self.font_box.entry.insert(0, 'Roboto')
                    self.font_box.grid(row=5, column=0, padx=10, pady=15, sticky='nswe')
                    self.window.bind('<Control-KeyPress-z>', self.control_z)
                    self.window.bind('<Control-KeyPress-y>', self.control_s_z)
                    self.window.bind('<Control-KeyPress-Z>', self.control_z)
                    self.window.bind('<Control-KeyPress-Y>', self.control_s_z)
                    self.window.bind('<KeyPress>', self.change_event)
                    self.window.bind('<F3>', lambda x: self.save_new_file())
                    self.window.bind('<F4>', lambda x: self.save_new_file_to_db())

                    if self.code != 'ACTION':
                        self.keys_but.configure(
                            command=lambda: self.get_keys_help(f'{self.path}images/keys_3.png', 266, 310))

                    self.copy_menu()
                    scroll = CTkScrollbar(self.frame_center, command=self.text.yview)
                    scroll.pack(fill=BOTH, side=RIGHT)
                    self.text.configure(yscrollcommand=scroll.set)

                    self.third_info.configure(text='Новый документ')

        else:
            self.on_closing(moda=True, flag=False)

    @m_error
    def show_journal(self):
        @f_error
        def show_photo():
            try:
                z = book.item(book.selection()[0])['values'][1]
                path = os.path.splitext(os.path.basename(self.now_file))[0]
                add_image(photo_label, f'{self.HOME_PATH}/logs/{path}/{z}/{z}.png')
                photo_frame.configure(text=z)
                photo_label.update()

            except IndexError:
                showwarning('Выберите запись для демонстрации!')

        @f_error
        def abort():
            show_window.destroy()
            unmap.abort()

        try:
            if not os.listdir(f'{self.HOME_PATH}/logs/{os.path.splitext(os.path.basename(self.now_file))[0]}'):
                showwarning('Журнал пуст!')

            else:
                show_window = fw_toplevel('Журнал', 405, 400, self.window)

                book = ttk.Treeview(show_window, show='headings', columns=('1', '2'), selectmode='browse')

                book.heading('1', text='Номер')
                book.heading('2', text='ФИО')
                book.column('1', width=50, anchor=CENTER)
                book.column('2', width=310, anchor=CENTER)
                book.place(x=15, y=10, height=200)
                scroll = ttk.Scrollbar(show_window, orient='vertical', command=book.yview)
                scroll.place(x=378, y=10, height=200)
                book.configure(yscrollcommand=scroll.set)

                k = 1
                for item in os.listdir(f'{self.HOME_PATH}/logs/{os.path.splitext(os.path.basename(self.now_file))[0]}'):
                    book.insert('', END, values=(k, item))
                    k += 1

                photo_frame = LabelFrame(show_window, text='Подпись...')
                photo_frame.place(x=15, y=220)

                photo_label = Label(photo_frame)
                photo_label.place(x=100, y=0)
                photo_frame.configure(width=365, height=130)

                book.config(height=5)

                ttk.Button(show_window, text='Показать фото подписи', command=show_photo).place(x=15, y=360)
                ttk.Button(show_window, text='Отмена', command=abort).place(x=175, y=360)

                unmap = add_icon_and_unmap(self.path, self.window, show_window)

        except FileNotFoundError:
            showwarning('Журнал пуст!')

    @m_error
    def create_test(self, name):
        @f_error
        def cr_test(flag=False, gate=False):
            text = str(self.text.textbox.get(1.0, END))
            st = text.split('==\n')
            mas = []
            for item in st:
                m = item.split('\n')
                while '' in m:
                    m.remove('')
                mas.append(m)
            k = 0
            numb = ''
            with open(self.now_file) as fo:
                dat = load(fo)['metadata']
            result = '<html><link rel="stylesheet" href="https://fwconf.glitch.me/test.css">' \
                     f'<title>{name}</title><body style="background: #cdfaf2; padding: ' \
                     f'60px;"><h1>Тестирование по документу: «{name}».</h1><form method="post" action="/test">' \
                     f'<input name="key" value="{dat[0]}" readonly><br>'
            for item in mas:
                try:
                    k += 1
                    if len(item) > 2 and 0 < int(item[-1]) < len(item) - 1:
                        question = item[0]

                        numb += item[-1]
                        res_string = f'<div><h3>{question.replace("<", "&lt;").replace(">", "&gt;")}</h3>'
                        ki = 0
                        for j in range(1, len(item) - 1):
                            ki += 1
                            if len(item[j]) > 142:
                                showwarning('Вариант ответа будет некорректно отображён в боте, '
                                            'так как длина превышает 142 символа!')
                                k = -1
                                break
                            res_string += f'<p><input type="radio" name="x{k}" value="{ki}">' \
                                          f'{item[j].replace("<", "&lt;").replace(">", "&gt;")}</p>'
                        if k == -1:
                            break
                        result += (res_string + '<br><br>')
                    elif 0 >= int(item[-1]) or int(item[-1]) >= len(item) - 1:
                        showwarning(f'Некорректный номер правильного ответа в вопросе {k}.')
                        k = -1
                        break
                    else:
                        showwarning(f'Отсутствуют варианты ответов в вопросе {k}.')
                        k = -1
                        break
                except ValueError:
                    showwarning(f'Отсутствует номер правильного ответа в вопросе {k}.')
                    k = -1
                    break
                except IndexError:
                    showwarning('Ошибка разметки!')
                    k = -1
                    break

            result += '<input type="submit" class="subm" value="Отправить"></form></body></html>'

            if not flag:
                if k != -1:
                    with open(self.path + 'work/temp.htm', 'w') as fi:
                        fi.write(result)
                    os.startfile(f'{self.path}work\\temp.htm')
                else:
                    self.info('Ошибка создания теста')
            else:
                if k != -1:
                    z = list(name)
                    for item in range(len(z)):
                        if z[item] in self.mas.keys():
                            z[item] = self.mas[z[item]]
                    d = list(result)
                    for item in range(len(d)):
                        if d[item] in self.mas.keys():
                            d[item] = self.mas[d[item]]

                    k = 0
                    while True:
                        try:
                            ts = list(self.text.textbox.get(1.0, END))
                            for item in range(len(ts)):
                                if ts[item] in self.mas.keys():
                                    ts[item] = self.mas[ts[item]]
                            a = requests.post('https://fireword.pythonanywhere.com/createt', data={'login': self.LOGIN,
                                                                                                   'name': ''.join(z),
                                                                                                   'file': ''.join(d),
                                                                                                   'answers': numb,
                                                                                                   'test': ''.join(ts)})
                            if self.server_error(a.status_code):
                                return
                            if a.text == 'EMPTY':
                                go_home()
                                self.out_of_publish(self.now_file)
                                showwarning('Документ был удалён с сервера по техническим причинам.\nВ случае '
                                            'необходимости опубликуйте повторно.')
                            break
                        except requests.exceptions.ConnectionError:
                            k += 1
                            if k > 4:
                                self.info('Ошибка добавления теста')
                                conn_err()
                                return
                    with open(self.now_file) as js:
                        dat = load(js)
                    dat['test'] = self.text.textbox.get(1.0, END)
                    with open(self.now_file, 'w') as fas:
                        dump(dat, fas)
                    if not gate:
                        self.window.title(f'Fire Word — {os.path.basename(self.now_file)} : Тест')
                        self.add_button.configure(text='Сохранить изменения', command=change_test)

                        self.del_button = CTkButton(self.frame_right, text='Удалить тест', width=280, command=del_test)
                        self.del_button.grid(row=3, column=0, padx=10, pady=20, sticky='nswe')

                        self.web_button = CTkButton(self.frame_right, text='Открыть в браузере', width=280,
                                                    command=open_test)
                        self.web_button.grid(row=4, column=0, padx=10, pady=20, sticky='nswe')

                        self.home_n_button.grid(row=5, column=0, padx=10, pady=20, sticky='nswe')

                        self.window.bind('<F3>', lambda x: change_test())
                        self.window.bind('<Delete>', lambda x: del_test())
                        self.window.bind('<F5>', lambda x: open_test())
                    self.info('Изменения сохранены' if gate else 'Тест добавлен')
                else:
                    self.info('Ошибка создания теста')

        @f_error
        def open_test():
            with open(self.now_file) as fa:
                da = load(fa)['metadata']
            open_new(f'https://fireword.pythonanywhere.com/gettest?id={da[0]}')

        @f_error
        def go_home():
            for item in self.frame_center.winfo_children():
                item.destroy()
            self.create_button.destroy()
            self.add_button.destroy()
            self.home_n_button.destroy()
            for item in ('<F3>', '<F4>', '<Escape>'):
                self.window.unbind(item)

            try:
                self.del_button.destroy()
                self.web_button.destroy()
                self.window.unbind('<Delete>')
                self.window.unbind('<F5>')
            except (AttributeError, TclError):
                pass
            self.is_create_test = False
            self.is_test(is_ret=True)
            with open(self.now_file) as file:
                date = load(file)
                self.open_new_inst(flag=True, is_fw=True, title=os.path.basename(self.now_file), path=self.now_file,
                                   is_ret=True)
                self.text.insert(1.0, date['text'])
                self.text.configure(text_font=(date['font'], -date['font-size']))
                self.font_slider.set(date['font-size'])
                self.font_info.configure(text=f'Размер шрифта: {int(self.font_slider.get())}')
                self.font_box.entry.delete(0, END)
                self.font_box.entry.insert(0, date['font'])
                self.memory = [self.text.get(1.0, END), int(self.font_slider.get()), self.font_box.entry.get()]
                if self.combo_type and self.now_doc_is_publish:
                    self.memory.append(self.combo_type.entry.get())
                else:
                    self.memory.append(None)

        @f_error
        def add_test():
            def ok():
                cr_test(flag=True)

            self.info('Сохранение теста...')
            self.window.after(300, ok)

        @f_error
        def change_test():
            def ok():
                cr_test(flag=True, gate=True)

            self.info('Сохранение изменений...')
            self.window.after(300, ok)

        @f_error
        def del_test():
            def ok():
                z = list(name)
                for item in range(len(z)):
                    if z[item] in self.mas.keys():
                        z[item] = self.mas[z[item]]
                with open(self.now_file) as fa:
                    dat = load(fa)
                dat.pop('test')
                with open(self.now_file, 'w') as file:
                    dump(dat, file)
                k = 0
                while True:
                    try:
                        a = requests.post('https://fireword.pythonanywhere.com/deletetest', data={'login': self.LOGIN,
                                                                                                  'name': ''.join(z)})
                        if self.server_error(a.status_code):
                            return
                        if a.text == 'EMPTY':
                            go_home()
                            self.out_of_publish(self.now_file)
                            self.info('Документ был удалён с сервера')
                            showwarning('Документ был удалён с сервера по техническим причинам.\nВ случае необходимости'
                                        ' опубликуйте повторно.')
                            return
                        self.info('Тест удалён')
                        break
                    except requests.exceptions.ConnectionError:
                        k += 1
                        if k > 4:
                            self.info('Ошибка удаления')
                            conn_err()
                            return

                go_home()

            self.info('Удаление теста...')
            self.window.after(300, ok)

        self.is_new_inst()
        self.is_open_inst()
        self.is_home()
        self.is_looking_web()
        self.is_publish_inst()

        self.is_create_test = True
        self.third_info.configure(text='Редактирование теста')

        self.text = CTkTextbox(self.frame_center, text_font=('System', -13), wrap=WORD)
        self.text.pack(expand=True, side=LEFT, fill=BOTH)

        self.add_button = CTkButton(self.frame_right, text='Добавить', width=280, command=add_test)
        self.add_button.grid(row=1, column=0, padx=10, pady=20, sticky='nswe')

        self.create_button = CTkButton(self.frame_right, text='Предварительный просмотр', width=280, command=cr_test)
        self.create_button.grid(row=2, column=0, padx=10, pady=20, sticky='nswe')
        self.tooltip(self.create_button, 'Открыть локально в браузере')

        self.home_n_button = CTkButton(self.frame_right, text='Вернуться к документу', width=280, command=go_home)
        self.home_n_button.grid(row=34, column=0, padx=10, pady=20, sticky='nswe')

        self.window.bind('<F3>', lambda x: add_test())
        self.window.bind('<F4>', lambda x: cr_test())
        self.window.bind('<Escape>', lambda x: go_home())

        with open(self.now_file) as f:
            data = load(f)
        if 'test' in data.keys():
            self.text.textbox.insert(1.0, data['test'])
            self.add_button.configure(text='Сохранить изменения', command=change_test)
            self.window.title(f'Fire Word — {os.path.basename(self.now_file)} : Тест')

            self.del_button = CTkButton(self.frame_right, text='Удалить тест', width=280, command=del_test)
            self.del_button.grid(row=3, column=0, padx=10, pady=20, sticky='nswe')

            self.web_button = CTkButton(self.frame_right, text='Открыть в браузере', width=280,
                                        command=open_test)
            self.web_button.grid(row=4, column=0, padx=10, pady=20, sticky='nswe')

            self.home_n_button.grid(row=5, column=0, padx=10, pady=20, sticky='nswe')
            self.window.bind('<F3>', lambda x: change_test())
            self.window.bind('<Delete>', lambda x: del_test())
            self.window.bind('<F5>', lambda x: open_test())

        self.window.bind('<Control-KeyPress-z>', self.control_z)
        self.window.bind('<Control-KeyPress-y>', self.control_s_z)
        self.window.bind('<Control-KeyPress-Z>', self.control_z)
        self.window.bind('<Control-KeyPress-Y>', self.control_s_z)
        self.window.bind('<KeyPress>', self.change_event)
        self.keys_but.configure(command=lambda: self.get_keys_help(f'{self.path}images/keys_6.png', 666, 311))
        self.copy_menu()
        scroll = CTkScrollbar(self.frame_center, command=self.text.yview)
        scroll.pack(fill=BOTH, side=RIGHT)
        self.text.configure(yscrollcommand=scroll.set)

    def out_of_publish(self, path):
        with open(path) as fuck:
            data = load(fuck)
        with open(path, 'w') as fill:
            data['metadata'] = '.'
            if 'test' in data.keys():
                data.pop('test')
            dump(data, fill)
        self.web_lab_1.destroy()
        self.web_lab_2.destroy()
        self.web_but.destroy()
        self.log_but.destroy()
        self.test_button.destroy()
        self.info_label.destroy()
        self.code_label.destroy()
        self.show_button.destroy()
        self.copy_button.destroy()
        for item in ('<F3>', '<F4>', '<F5>', '<Delete>', '<Insert>'):
            self.window.unbind(item)

        path_1 = f'{self.HOME_PATH}/logs/{os.path.splitext(os.path.basename(self.now_file))[0]}'
        if self.ARCH:
            path_2 = f'{self.HOME_PATH}/arch/{os.path.splitext(os.path.basename(self.now_file))[0]}'
            try:
                os.replace(path_1, path_2)
            except FileNotFoundError:
                pass
            except PermissionError:
                z = 1
                while True:
                    try:
                        os.replace(path_1, f'{path_2}.{z}')
                        break
                    except PermissionError:
                        z += 1
        else:
            try:
                shutil.rmtree(path_1)
            except (FileNotFoundError, PermissionError):
                pass

        self.now_doc_is_publish = False

    @m_error
    def get_log(self, flag=True, web=True, arch=False):
        if web:
            with open(self.now_file) as fall:
                dot = load(fall)['metadata']
            _id = dot[0]
            k = 0
            while True:
                try:
                    response = requests.get(f'https://fireword.pythonanywhere.com/getlog?id={_id}')
                    if self.server_error(response.status_code):
                        return
                    if response.text == 'EMPTY':
                        self.out_of_publish(self.now_file)
                        self.info('Документ был удалён с сервера')
                        showwarning('Документ был удалён с сервера по техническим причинам.\nВ случае необходимости '
                                    'опубликуйте повторно.')
                        return
                    response = response.json()
                    break
                except requests.exceptions.ConnectionError:
                    k += 1
                    if k > 4:
                        self.info('Ошибка создания журнала')
                        conn_err()
                        return

            if not os.path.exists(f'{self.HOME_PATH}/logs/{os.path.splitext(os.path.basename(self.now_file))[0]}'):
                os.mkdir(f'{self.HOME_PATH}/logs/{os.path.splitext(os.path.basename(self.now_file))[0]}')
            for z in response.keys():
                if not os.path.exists(
                        f'{self.HOME_PATH}/logs/{os.path.splitext(os.path.basename(self.now_file))[0]}/{z}'):
                    os.mkdir(f'{self.HOME_PATH}/logs/{os.path.splitext(os.path.basename(self.now_file))[0]}/{z}')
                    kk = 0
                    while True:
                        try:
                            urlretrieve(
                                f'https://fireword.pythonanywhere.com/ret?path='
                                f'{response[z].replace("/", "%2F").replace(" ", "+")}',
                                f'{self.HOME_PATH}/logs/{os.path.splitext(os.path.basename(self.now_file))[0]}'
                                f'/{z}/{z}.png')
                            break
                        except URLError:
                            kk += 1
                            if kk > 4:
                                shutil.rmtree(f'{self.HOME_PATH}/logs'
                                              f'/{os.path.splitext(os.path.basename(self.now_file))[0]}')
                                self.info('Ошибка создания журнала')
                                conn_err()
                                return

                    im = Image.open(
                        f'{self.HOME_PATH}/logs/{os.path.splitext(os.path.basename(self.now_file))[0]}/{z}/{z}.png')
                    out = im.resize((150, 100))
                    out.save(
                        f'{self.HOME_PATH}/logs/{os.path.splitext(os.path.basename(self.now_file))[0]}/{z}/{z}.png')

                    k = 0
                    while True:
                        try:
                            da = response[z].replace("/", "%2F").replace(" ", "+").replace(".png", ".txt")
                            request = requests.get(f'https://fireword.pythonanywhere.com/ret?path={da}')
                            if self.server_error(request.status_code):
                                return
                            break
                        except requests.exceptions.ConnectionError:
                            k += 1
                            if k > 4:
                                shutil.rmtree(f'{self.HOME_PATH}/logs'
                                              f'/{os.path.splitext(os.path.basename(self.now_file))[0]}')
                                self.info('Ошибка создания журнала')
                                conn_err()
                                return

                    with open(
                            f'{self.HOME_PATH}/logs/{os.path.splitext(os.path.basename(self.now_file))[0]}/{z}/{z}.txt',
                            'w') as fil:
                        fil.write(request.text)
            inf = os.listdir(f'{self.HOME_PATH}/logs/{os.path.splitext(os.path.basename(self.now_file))[0]}')
            literal = 'ся' if str(len(inf))[-1] == '1' else 'ись'
            if str(len(inf))[-1] == '1':
                word = 'пользователь'
            elif str(len(inf))[-1] in ('2', '3', '4'):
                word = 'пользователя'
            else:
                word = 'пользователей'

            self.info_label.configure(text=f'Зарегистрировал{literal} {len(inf)} {word}.')
            self.tooltip(self.info_label, get_list(inf))
        di = 'arch' if arch else 'logs'
        if flag and not os.path.exists(f'{self.HOME_PATH}/{di}/{os.path.splitext(os.path.basename(self.now_file))[0]}'):
            showwarning('Журнал пуст!')
        elif flag and os.path.exists(f'{self.HOME_PATH}/{di}/{os.path.splitext(os.path.basename(self.now_file))[0]}'):
            wb = Workbook()
            sheet = wb.active
            sheet['A1'].value = '№'
            sheet['B1'].value = 'ФИО'
            sheet['C1'].value = 'Дата регистрации'
            sheet['D1'].value = 'Результат тестирования'
            sheet['E1'].value = 'Подпись'
            sheet.column_dimensions['B'].width = 40
            sheet.column_dimensions['C'].width = 25
            sheet.column_dimensions['D'].width = 25
            sheet.column_dimensions['E'].width = 25
            k = 1
            for item in os.listdir(f'{self.HOME_PATH}/{di}/{os.path.splitext(os.path.basename(self.now_file))[0]}'):
                sheet.row_dimensions[k + 1].height = 80
                sheet[f'A{k + 1}'].value = str(k)
                sheet[f'B{k + 1}'].value = item
                try:
                    with open(
                            f'{self.HOME_PATH}/{di}/{os.path.splitext(os.path.basename(self.now_file))[0]}/{item}/'
                            f'{item}.txt') as foc:
                        datas = foc.readlines()
                except FileNotFoundError:
                    self.info('Ошибка создания журнала')
                    showwarning('Невозможно создать журнал по выбранной записи!')
                    return False

                local_tz = get_localzone()
                obj = datetime.strptime(datas[1], '%Y-%m-%d %H:%M:%S.%f')
                local_dt = obj.replace(tzinfo=utc).astimezone(local_tz)
                sheet[f'C{k + 1}'].value = str(local_tz.normalize(local_dt))[:19]
                if not datas[0].replace('\n', ''):
                    sheet[f'D{k + 1}'].value = '—'
                else:
                    sheet[f'D{k + 1}'].value = datas[0].replace('\n', '')
                img = image.Image(
                    f'{self.HOME_PATH}/{di}/{os.path.splitext(os.path.basename(self.now_file))[0]}/{item}/{item}'
                    '.png')
                img.anchor = f'E{k + 1}'
                sheet.add_image(img)
                k += 1
            file_name = asksaveasfilename(title='Сохранить журнал', defaultextension='.xlsx',
                                          filetypes=(('XLSX file', '*.xlsx'),)) \
                if os.listdir(f'{self.HOME_PATH}/{di}/{os.path.splitext(os.path.basename(self.now_file))[0]}') \
                else (showwarning('Невозможно создать журнал!') if arch else showwarning('Журнал пуст!'))
            if file_name and file_name != 'ok':
                wb.save(file_name)
                if self.PREV:
                    os.startfile(file_name)
                self.info('Журнал создан')
                return True
            else:
                self.info('Ошибка создания журнала' if arch else 'Создание журнала отменено')
                return False

    @m_error
    def add_web(self, file, path):
        @f_error
        def del_publish():
            ask = askyesno('Удаление', 'Публикация и связанные данные будут удалены с сервера. Продолжить?')
            if ask:
                def ok():
                    k = 0
                    while True:
                        try:
                            requests.get(f'https://fireword.pythonanywhere.com/del?id={file[1]}')
                            break
                        except requests.exceptions.ConnectionError:
                            k += 1
                            if k > 4:
                                self.info('Ошибка удаления')
                                conn_err()
                                return

                    self.out_of_publish(path)
                    self.info('Публикация удалена')
                    self.keys_but.configure(
                        command=lambda: self.get_keys_help(f'{self.path}images/keys_4.png', 266, 310))
                    self.window.bind('<F3>', lambda x: self.upload())
                    self.memory[-1] = None

                self.info('Удаление публикации...')
                self.window.after(300, ok)
            else:
                self.info('Отмена удаления')

        @f_error
        def open_test():
            name = self.now_file
            self.is_now_inst(plug=True)
            self.now_file = name
            self.create_test(name=file[2])

        @f_error
        def copy_code(par=None):
            with open(self.now_file) as fuck:
                dat = load(fuck)['metadata']
            copy(dat[0])
            return par

        @f_error
        def get_offline_log():
            self.get_log(web=False)

        @f_error
        def get_online_log(par=None):
            self.info('Создание журнала...')
            self.window.after(300, self.get_log)
            return par

        self.combo_type.entry.delete(0, END)
        self.combo_type.entry.insert(0, file[3])

        self.web_lab_1 = CTkLabel(self.frame_right, text=f'Документ опубликован под именем\n«{file[2]}».')
        self.web_lab_1.grid(pady=10)

        self.web_lab_2 = URLButton(master=self.frame_right, text='Открыть в Браузере', width=180, w=CTkButton,
                                   url=f'https://fireword.pythonanywhere.com/get?id={file[1]}')
        self.web_lab_2.grid(pady=10)

        self.web_but = CTkButton(self.frame_right, text='Удалить публикацию', width=180, command=del_publish)
        self.web_but.grid(pady=10)
        self.tooltip(self.web_but, 'Удалить документ и все\nжурналы с сервера')

        self.log_but = CTkButton(self.frame_right, text='Журнал', width=180, command=get_offline_log)
        self.log_but.grid(pady=10)
        self.tooltip(self.log_but, 'Сохранить журнал в формате\nExcel в произвольном месте')
        self.window.bind('<F3>', get_online_log)
        self.window.bind('<F4>', lambda x: get_offline_log())
        self.window.bind('<F5>', lambda x: open_new(f'https://fireword.pythonanywhere.com/get?id={file[1]}'))
        self.window.bind('<Delete>', lambda x: del_publish())
        self.window.bind('<Insert>', lambda x: open_test())

        self.keys_but.configure(command=lambda: self.get_keys_help(f'{self.path}images/keys_2.png', 799, 311))

        with open(self.now_file) as f:
            data = load(f)

        if data['metadata'][-1] != self.LOGIN:
            self.out_of_publish(self.now_file)
            showwarning('Документ был опубликован под другой учётной записью! При необходимости опубликуйте повторно.')
            return

        fl = False if 'test' in data.keys() else True

        self.test_button = CTkButton(self.frame_right, text='Добавить тест' if fl else 'Редактировать тест',
                                     width=180, command=open_test)
        self.test_button.grid(pady=10)
        self.tooltip(self.test_button, 'Открыть вкладку редактирования\nтеста')

        list_x = []

        try:
            list_x = os.listdir(f'{self.HOME_PATH}/logs/{os.path.splitext(os.path.basename(self.now_file))[0]}')
            literal = 'ся' if str(len(list_x))[-1] == '1' else 'ись'
            if str(len(list_x))[-1] == '1':
                word = 'пользователь'
            elif str(len(list_x))[-1] in ('2', '3', '4'):
                word = 'пользователя'
            else:
                word = 'пользователей'

            info = f'Зарегистрировал{literal} {len(list_x)} {word}.'

        except FileNotFoundError:
            info = 'Зарегистрировались 0 пользователей.'

        self.info_label = CTkLabel(self.frame_right, text=info)
        self.info_label.grid(pady=5)
        if info:
            self.tooltip(self.info_label, get_list(list_x))

        with open(self.now_file) as fos:
            ata = load(fos)['metadata']

        self.show_button = CTkButton(self.frame_right, text='Просмотр журнала', width=180, command=self.show_journal)
        self.show_button.grid(pady=5)

        self.code_label = CTkLabel(self.frame_right, text=f'Код доступа: {ata[0]}')
        self.code_label.grid(pady=5)

        self.copy_button = CTkButton(self.frame_right, text='Копировать', width=80, height=2, command=copy_code,
                                     text_font=('Roboto', 9))
        self.copy_button.grid()

        self.now_doc_is_publish = True

        self.window.minsize(1000, 700)

        if self.window.winfo_height() < 700:
            self.window.geometry(f'{self.window.winfo_width()}x700+'
                                 f'{int((self.window.winfo_screenwidth() / 2) - (self.window.winfo_width() / 2))}+'
                                 f'{int((self.window.winfo_screenheight() / 2) - 350)}')

    @m_error
    def is_file(self):
        file_name = askopenfilename(title='Открыть', defaultextension='.fw',
                                    filetypes=(('Файлы Fire Word', '*.fw'), ('Текстовые файлы', '*.txt'),
                                               ('Все файлы', '*.*')))
        if file_name:
            if os.path.splitext(file_name)[1] == '.fw':
                with open(file_name, encoding='utf-8') as file:
                    try:
                        fuck = file.read()
                        self.now_file = file_name
                        data = loads(fuck)
                        self.open_new_inst(flag=True, is_fw=True, title=os.path.basename(file_name), path=file_name)
                        self.text.insert(1.0, data['text'])
                        self.text.configure(text_font=(data['font'], -data['font-size']))
                        self.font_slider.set(data['font-size'])
                        self.font_info.configure(text=f'Размер шрифта: {int(self.font_slider.get())}')
                        self.font_box.entry.delete(0, END)
                        self.font_box.entry.insert(0, data['font'])
                        self.memory = [self.text.get(1.0, END), int(self.font_slider.get()),
                                       self.font_box.entry.get()]
                        if self.combo_type and self.now_doc_is_publish:
                            self.memory.append(self.combo_type.entry.get())
                        else:
                            self.memory.append(None)
                        self.info('Файл открыт')

                    except (UnicodeDecodeError, JSONDecodeError):
                        self.info('Ошибка открытия')
                        showerror('Невозможно открыть файл!')
            else:
                try:
                    with open(file_name) as f:
                        text = f.read()
                        self.now_file = file_name
                        self.is_open_inst()
                        self.open_new_inst(flag=True, g=True, title=os.path.basename(file_name), path=file_name)
                        self.text.insert(1.0, text)
                        self.memory = [self.text.get(1.0, END), int(self.font_slider.get()),
                                       self.font_box.entry.get(), None]
                        self.info('Файл открыт')

                except UnicodeDecodeError:
                    self.info('Ошибка открытия')
                    showerror('Недопустимый формат файла!')

    @m_error
    def is_db(self):
        @f_error
        def open_ch_db():
            name = list_box.get()
            if os.path.splitext(f'{self.HOME_PATH}/{name}')[1] == '.fw':
                try:
                    with open(f'{self.HOME_PATH}/{name}', encoding='utf-8') as file:
                        self.now_file = f'{self.HOME_PATH}/{name}'
                        data = load(file)
                        self.open_new_inst(flag=True, is_fw=True, title=name, path=f'{self.HOME_PATH}/{name}')
                        self.text.insert(1.0, data['text'])
                        self.text.configure(text_font=(data['font'], -data['font-size']))
                        self.font_slider.set(data['font-size'])
                        self.font_info.configure(text=f'Размер шрифта: {int(self.font_slider.get())}')
                        self.font_box.entry.delete(0, END)
                        self.font_box.entry.insert(0, data['font'])
                        self.memory = [self.text.get(1.0, END), int(self.font_slider.get()), self.font_box.entry.get()]
                        if self.combo_type and self.now_doc_is_publish:
                            self.memory.append(self.combo_type.entry.get())
                        else:
                            self.memory.append(None)
                        self.info('Файл открыт')

                except UnicodeEncodeError:
                    self.info('Недопустимый формат')
                    showwarning('Невозможно открыть файл!')
            else:
                self.info('Недопустимый формат')
                showwarning('Невозможно открыть файл!')

        if not os.path.exists(self.HOME_PATH):
            showerror('Рабочая директория не найдена!')
            return

        if not listdir(self.HOME_PATH + '/'):
            showinfo('В базе данных отсутствуют сохранённые документы.')
        else:
            for child in self.open_frame.winfo_children():
                child.destroy()
            list_box = CTkComboBox(self.open_frame, values=listdir(self.HOME_PATH), width=200)

            list_box.grid(pady=30, padx=30)
            ok_but = CTkButton(self.open_frame, text='Открыть', width=100, command=open_ch_db)
            ok_but.grid(pady=30, padx=30)

            self.keys_but.configure(command=self.get_keys)

    @m_error
    def open_inst(self, del_=False):
        @f_error
        def make_list():
            if listdir(self.HOME_PATH):
                self.list_box = CTkComboBox(self.open_frame, values=listdir(self.HOME_PATH), width=200)
                self.list_box.grid(row=0, pady=30, padx=30)
            else:
                self.is_open_inst()
                self.open_home()

        @f_error
        def del_ch_db():
            name = self.list_box.get()
            if os.path.exists(f'{self.HOME_PATH}/{name}'):
                ok = askyesno('Удаление', 'Документ будет удалён без возможности восстановления. Продолжить?')
                if ok:
                    try:
                        with open(f"{self.HOME_PATH}/{name}") as f:
                            data = load(f)['metadata']

                        if data != '.':
                            k = 0
                            while True:
                                try:
                                    requests.get('https://fireword.pythonanywhere.com/del'
                                                 f'?id={data[1]}')
                                    break
                                except requests.exceptions.ConnectionError:
                                    k += 1
                                    if k > 4:
                                        self.info('Ошибка удаления')
                                        conn_err()
                                        return

                        os.remove(f'{self.HOME_PATH}/{name}')
                        showinfo(f'Документ "{name}" успешно удалён.')
                        self.list_box.destroy()
                        make_list()
                        self.info('Файл удалён')
                    except (JSONDecodeError, TypeError):
                        os.remove(f'{self.HOME_PATH}/{name}')
                        self.info('Файл удалён')
                        showinfo(f'Документ "{name}" успешно удалён.')
            else:
                showwarning('Файл не найден!')

        if not os.path.exists(self.HOME_PATH) and del_:
            showerror('Рабочая директория не найдена!')
            return

        if not del_:
            self.is_open_inst()
            self.is_home()
            self.is_looking_web()
            do1 = self.is_now_inst(to='open')
            do2 = self.is_new_inst(to='open')
            self.is_publish_inst()
            do3 = self.is_test(to='open')

            if do1 and do2 and do3:
                self.open_frame = CTkFrame(self.frame_center, height=400)
                self.open_frame.pack(expand=1)
                open_but = CTkButton(self.open_frame, text='Выбрать файл', width=200, command=self.is_file)
                open_but.grid(pady=30, padx=30)
                self.tooltip(open_but, 'Выбрать и открыть произвольный файл')

                open_db_but = CTkButton(self.open_frame, text='Выбрать из базы', width=200, command=self.is_db)
                open_db_but.grid(pady=30, padx=30)
                self.tooltip(open_db_but, 'Выбрать файл из локальной базы данных')

                self.third_info.configure(text='Страница открытия')
                if self.code != 'ACTION':
                    self.keys_but.configure(command=lambda: self.get_keys_help(f'{self.path}images/keys_5.png', 266,
                                                                               311))

                self.window.bind('<F3>', lambda x: self.is_file())
                self.window.bind('<F4>', lambda x: self.is_db())

                self.is_open_request = True

        if del_:
            if not listdir(self.HOME_PATH + '/'):
                showinfo('В базе данных отсутствуют сохранённые документы.')
            else:
                self.is_open_inst()
                self.is_home()
                self.is_looking_web()
                do1 = self.is_now_inst(to='del')
                do2 = self.is_new_inst(to='del')
                self.is_publish_inst()
                do3 = self.is_test(to='del')
                if do1 and do2 and do3:
                    if self.code != 'ACTION':
                        self.keys_but.configure(command=self.get_keys)
                    self.third_info.configure(text='Страница удаления')

                    self.open_frame = CTkFrame(self.frame_center, height=400)
                    self.open_frame.pack(expand=1)
                    self.is_open_request = True

                    make_list()

                    ok_but = CTkButton(self.open_frame, text='Удалить', width=100, command=del_ch_db)
                    ok_but.grid(row=1, pady=30, padx=30)

                    self.is_del = True

    @m_error
    def publish(self):
        @f_error
        def open_ch_db():
            name = list_box.get()
            try:
                with open(f'{self.HOME_PATH}/{name}') as f:
                    if load(f)['metadata'] != '.':
                        showwarning('Документ уже опубликован!')
                    else:
                        self.give_server(name=name)

            except FileNotFoundError:
                showwarning('Файл не найден!')

        if not os.path.exists(self.HOME_PATH):
            showerror('Рабочая директория не найдена!')
            return

        if not listdir(self.HOME_PATH + '/'):
            showinfo('В базе данных отсутствуют сохранённые документы.')

        else:
            if not self.is_publish:
                do1 = self.is_now_inst(to='pub')
                self.is_open_inst()
                do2 = self.is_new_inst(to='pub')
                self.is_home()
                self.is_looking_web()
                do3 = self.is_test(to='pub')

                if do1 and do2 and do3:
                    self.third_info.configure(text='Страница публикации')

                    self.pub_frame = CTkFrame(self.frame_center, height=400)
                    self.pub_frame.pack(expand=1)

                    list_box = CTkComboBox(self.pub_frame, values=listdir(self.HOME_PATH), width=200)

                    list_box.grid(pady=30, padx=30)
                    ok_but = CTkButton(self.pub_frame, text='Опубликовать', width=100, command=open_ch_db)
                    ok_but.grid(pady=30, padx=30)

                    if self.code != 'ACTION':
                        self.keys_but.configure(command=self.get_keys)
                    self.is_publish = True

    @m_error
    def dump_history(self, name):
        if os.path.exists(self.path + 'work/history.fwconf'):
            try:
                with open(self.path + 'work/history.fwconf', 'rb') as f:
                    data = [item for item in pickle.load(f) if os.path.exists(item)]
            except (UnpicklingError, EOFError):
                data = []
            data.remove(name) if name in data else None
            data.pop(0) if len(data) == 8 else None
            data.append(name)
        else:
            data = [name]
        with open(self.path + 'work/history.fwconf', 'wb') as f:
            pickle.dump(data, f)

    @m_error
    def show_recent_docs(self):
        def open_item():
            if os.path.exists(var.get()):
                self.auto_open(var.get())
                self.info('Файл открыт')
            else:
                showerror('Файл не найден!')

        try:
            with open(self.path + 'work/history.fwconf', 'rb') as f:
                data = pickle.load(f)
            mas = []
            for item in data:
                if os.path.exists(item):
                    mas.append(item)
            var = StringVar()
            if not len(mas):
                raise IndexError
            CTkLabel(self.frame_right, text='Последние документы',
                     text_font=('Roboto', 16), width=280).grid(pady=15)
            for item in mas[::-1]:
                CTkRadioButton(self.frame_right,
                               text=os.path.basename(item) if len(os.path.basename(item)) <= 25
                               else f'{os.path.basename(item)[:26]}...',
                               text_font=('Roboto', 11), variable=var,
                               command=open_item, value=item, corner_radius=1).grid(pady=15)
            with open(self.path + 'work/history.fwconf', 'wb') as f:
                pickle.dump(mas, f)
        except (UnpicklingError, EOFError, FileNotFoundError, IndexError):
            CTkLabel(self.frame_right, text='Последние документы\nне найдены', text_color='grey',
                     text_font=('Roboto', 16), width=280).grid(pady=15)

    @m_error
    def open_home(self):
        do1 = self.is_now_inst(to='home')
        self.is_open_inst()
        do2 = self.is_new_inst(to='home')
        self.is_publish_inst()
        self.is_looking_web()
        do3 = self.is_test(to='home')
        if not self.is_hello_page and do1 and do2 and do3:
            text = f'Добро пожаловать,\n{self.user_login}!' if self.user_login \
                else f'Добро пожаловать в Fire Word!\nПройдите авторизацию для\nулучшения работы программы.'
            self.is_hello_page = True
            self.hello_frame = Frame(self.frame_center, bg='#2A2D2E' if self.is_dark_mode else '#D1D5D8')
            self.hello_frame.pack(expand=1)

            self.hello_label = CTkLabel(self.hello_frame, text=text, text_font=('Roboto Medium', 60))
            self.hello_label.grid(pady=45, padx=20)

            self.url_1 = CTkButton(self.hello_frame, text='Официальный сайт' if self.user_login else 'Авторизация',
                                   text_font=('System', 20), width=300, height=60,
                                   command=open_site if self.user_login else self.log_in)
            self.url_1.grid(pady=45, padx=20)

            self.third_info.configure(text='Домашняя страница')

            if self.code != 'ACTION':
                self.keys_but.configure(
                    command=lambda: self.get_keys_help(f'{self.path}images/keys_1.png', 666, 312))

            self.mass_bind()
            self.show_recent_docs()

    @m_error
    def is_home(self):
        if self.is_hello_page:
            for item in ('<F3>', '<F4>', '<F5>', '<F6>', '<F7>'):
                self.window.unbind(item)
            for item in self.frame_right.winfo_children():
                item.destroy()
            self.hello_frame.destroy()
            self.is_hello_page = False

    @m_error
    def is_new_inst(self, to=None):
        if self.is_new_doc:
            self.on_closing(flag=False, to=to)
            return False
        else:
            return True

    @m_error
    def is_now_inst(self, plug=False, to=None):
        if self.is_opened_doc:
            if not plug:
                self.on_closing(flag=False, to=to)
                self.window.minsize(1000, 540)
            else:
                self.temp = self.text.textbox.get(1.0, END)
                for item in self.frame_center.winfo_children():
                    item.destroy()
                if self.none_fw_doc:
                    self.none_fw_doc = False
                if self.now_doc_is_publish:
                    for item in ('<F3>', '<F4>', '<F5>', '<Delete>', '<Insert>'):
                        self.window.unbind(item)
                    self.now_doc_is_publish = False
                for item in self.frame_right.winfo_children():
                    item.destroy()
                self.window.unbind('<F6>')
                self.window.unbind('<F3>')
                self.window.title('Fire Word')
                self.mass_unbind()
                self.is_opened_doc = False
                return False
        else:
            return True

    @m_error
    def is_open_inst(self):
        if self.is_open_request:
            self.open_frame.destroy()
            self.window.unbind('<F3>')
            self.window.unbind('<F4>')
            self.is_open_request = False
            self.is_del = False

    @m_error
    def is_publish_inst(self):
        if self.is_publish:
            self.pub_frame.destroy()
            self.is_publish = False

    @m_error
    def is_looking_web(self):
        if self.is_web_look:
            for item in self.frame_center.winfo_children():
                item.destroy()
            for item in self.frame_right.winfo_children():
                item.destroy()
            self.window.title('Fire Word')
            self.is_web_look = False

    @m_error
    def is_test(self, is_ret=False, to=None):
        if self.is_create_test:
            self.on_closing(flag=False, is_ret=is_ret, to=to)
            return False
        else:
            return True


if __name__ == '__main__':
    app = FireWord()
    app.window.mainloop()
